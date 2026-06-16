from __future__ import annotations

import glob
import inspect
import json
import os
import re
import shutil
import subprocess
import tempfile
import time
from collections import defaultdict
from dataclasses import asdict, dataclass, field
from datetime import datetime
from multiprocessing import Process
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    import pandas as pd
except ImportError:
    pd = None

try:
    import pyodbc
except ImportError:
    pyodbc = None

try:
    from docx import Document
    from docx.shared import RGBColor
except ImportError:
    Document = None
    RGBColor = None

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
except ImportError:
    SimpleDocTemplate = None
try:
    from PIL import Image, ImageTk
except ImportError:
    Image = None
    ImageTk = None

APP_FONT = ("Arial", 10)
DEFAULT_TRACE_DIR = r"C:\Program Files (x86)\HAMILTON\LogFiles"
DEFAULT_RESULTS_DIR = r"C:\Results\Sequence Analysis"

ROLE_OPTIONS = ["tip pick up", "aspirate", "dispense"]
TRANSFER_OPTIONS = ["samples", "buffer", "STD/QCs"]
REPLICATE_OPTIONS = ["single", "duplicate", "triplicate"]
REPLICATE_MAP = {"single": 1, "duplicate": 2, "triplicate": 3}
UI_BG = "#000000"
UI_PANEL_BG = "#000000"
UI_TEXT_BG = "#2b004f"
UI_TEXT_FG = "white"
UI_BUTTON_BG = "#2b004f"
UI_BUTTON_ACTIVE = "#4b0082"
CHECK_SELECT_BG = "white"
CHANNEL_MODE_OPTIONS = ["channels", "head96", "head384"]
VALID_CHANNEL_PATTERN_LENGTHS = {8, 12, 14, 16}

PYHAMILTON_IMPORT_ERROR = ""

try:
    import pyhamilton
    import pyhamilton.interface as phi

    from pyhamilton import (
        HamiltonError,
        HamiltonStepError,
        HamiltonTimeoutError,
        LayoutManager,
        OEM_HSL_PATH,
        OEM_RUN_EXE_PATH,
        Plate96,
        ResourceType,
        initialize,
        normal_logging,
    )

except Exception as exc:
    PYHAMILTON_IMPORT_ERROR = f"{type(exc).__name__}: {exc}"
    pyhamilton = None
    phi = None
    HamiltonStepError = Exception
    HamiltonError = Exception
    HamiltonTimeoutError = Exception
    LayoutManager = None
    Plate96 = None
    ResourceType = None
    OEM_RUN_EXE_PATH = ""
    OEM_HSL_PATH = ""


def _import_optional_pyhamilton_module(module_name: str):
    try:
        return __import__(module_name, fromlist=["*"])
    except Exception:
        return None


PYHAMILTON_MODULES = []

if pyhamilton is not None:
    PYHAMILTON_MODULES.extend(
        module
        for module in [
            pyhamilton,
            phi,
            _import_optional_pyhamilton_module("pyhamilton.liquid_handling_wrappers"),
            _import_optional_pyhamilton_module("pyhamilton.defaultcmds"),
            _import_optional_pyhamilton_module("pyhamilton.oemerr"),
            _import_optional_pyhamilton_module("pyhamilton.interface"),
        ]
        if module is not None
    )


def _get_optional_pyhamilton_function(*names):
    for module in PYHAMILTON_MODULES:
        for name in names:
            if hasattr(module, name):
                return getattr(module, name)
    return None


def _call_best_effort(func, *args, **kwargs):
    if func is None:
        raise RuntimeError("PyHamilton wrapper function is unavailable.")

    try:
        signature = inspect.signature(func)
    except Exception:
        signature = None

    if signature is None:
        return func(*args, **kwargs)

    allowed_kwargs = {
        key: value
        for key, value in kwargs.items()
        if key in signature.parameters
    }

    try:
        return func(*args, **allowed_kwargs)
    except TypeError:
        return func(*args)


def _send_command_best_effort(ham_int, command_name: str, **kwargs):
    last_error = None

    command_candidates = [
        command_name,
        command_name.upper(),
        command_name.lower(),
    ]

    for module in PYHAMILTON_MODULES:
        for candidate in command_candidates:
            if hasattr(module, candidate):
                command_candidates.append(getattr(module, candidate))

    seen = []
    deduped_candidates = []

    for candidate in command_candidates:
        key = str(candidate)

        if key not in seen:
            seen.append(key)
            deduped_candidates.append(candidate)

    if hasattr(ham_int, "send_command"):
        for candidate in deduped_candidates:
            try:
                tid = ham_int.send_command(command=candidate, **kwargs)

                if hasattr(ham_int, "wait_on_response"):
                    return ham_int.wait_on_response(
                        tid,
                        raise_first_exception=True,
                        timeout=120,
                    )

                return tid
            except Exception as exc:
                last_error = exc

            try:
                tid = ham_int.send_command(candidate, **kwargs)

                if hasattr(ham_int, "wait_on_response"):
                    return ham_int.wait_on_response(
                        tid,
                        raise_first_exception=True,
                        timeout=120,
                    )

                return tid
            except Exception as exc:
                last_error = exc

    raise RuntimeError(f"Unable to send PyHamilton command {command_name}: {last_error}")


def _clean_liquid_class(liq_class: str) -> str:
    return str(liq_class or "").strip()


def ph_tip_pick_up_seq(ham_int, tipseq: str, channel: str):
    channel = normalize_assigned_channel_pattern(channel, "channels")

    func = _get_optional_pyhamilton_function(
        "tip_pick_up_seq",
        "tip_pickup_seq",
        "tip_pick_up_seq2",
        "tip_pick_up",
        "tip_pickup",
    )

    if func:
        return _call_best_effort(
            func,
            ham_int,
            tipseq=tipseq,
            tip_seq=tipseq,
            sequence=tipseq,
            channel=channel,
            channels=channel,
            ch_patt=channel,
        )

    return _send_command_best_effort(
        ham_int,
        "PICKUP",
        tipSequence=tipseq,
        channelVariable=channel,
        sequenceCounting=1,
    )


def ph_aspirate_seq(ham_int, asp_seq: str, vols: float, channel: str, liq_class: str):
    channel = normalize_assigned_channel_pattern(channel, "channels")
    liq_class = _clean_liquid_class(liq_class)

    print(f"USING CHANNEL ASPIRATE PATTERN: {channel}")
    print(f"USING CHANNEL ASPIRATE LIQUID CLASS: {liq_class}")

    func = _get_optional_pyhamilton_function(
        "aspirate_seq",
        "aspirate_from_seq",
        "aspirate_ch_seq",
        "aspirate",
    )

    if func:
        return _call_best_effort(
            func,
            ham_int,
            asp_seq=asp_seq,
            sequence=asp_seq,
            vols=vols,
            volume=vols,
            channel=channel,
            channels=channel,
            ch_patt=channel,
            liq_class=liq_class,
            liquid_class=liq_class,
            lc=liq_class,
        )

    return _send_command_best_effort(
        ham_int,
        "ASPIRATE",
        aspirateSequence=asp_seq,
        labwarePositions="",
        volumes=vols,
        channelVariable=channel,
        liquidClass=liq_class,
        sequenceCounting=0,
        capacitiveLLD=2,
    )


def ph_dispense_seq(ham_int, disp_seq: str, vols: float, channel: str, liq_class: str):
    channel = normalize_assigned_channel_pattern(channel, "channels")
    liq_class = _clean_liquid_class(liq_class)

    print(f"USING CHANNEL DISPENSE PATTERN: {channel}")
    print(f"USING CHANNEL DISPENSE LIQUID CLASS: {liq_class}")

    func = _get_optional_pyhamilton_function(
        "dispense_seq",
        "dispense_to_seq",
        "dispense_seq_countoff",
        "dispense",
    )

    if func:
        return _call_best_effort(
            func,
            ham_int,
            disp_seq=disp_seq,
            sequence=disp_seq,
            vols=vols,
            volume=vols,
            channel=channel,
            channels=channel,
            ch_patt=channel,
            liq_class=liq_class,
            liquid_class=liq_class,
            lc=liq_class,
        )

    return _send_command_best_effort(
        ham_int,
        "DISPENSE",
        dispenseSequence=disp_seq,
        labwarePositions="",
        volumes=vols,
        channelVariable=channel,
        liquidClass=liq_class,
        sequenceCounting=0,
        channelUse=1,
    )


def ph_tip_eject_seq2(ham_int, waste_seq: str, channel: str):
    channel = normalize_assigned_channel_pattern(channel, "channels")

    func = _get_optional_pyhamilton_function(
        "tip_eject_seq2",
        "tip_eject_seq",
        "tip_eject3",
        "tip_eject",
    )

    if func:
        return _call_best_effort(
            func,
            ham_int,
            waste_seq=waste_seq,
            sequence=waste_seq,
            channel=channel,
            channels=channel,
            ch_patt=channel,
        )

    return _send_command_best_effort(
        ham_int,
        "EJECT",
        wasteSequence=waste_seq,
        labwarePositions="",
        channelVariable=channel,
        sequenceCounting=0,
    )


def ph_inc_sequence(ham_int, sequence: str, increment: int):
    func = _get_optional_pyhamilton_function(
        "inc_sequence",
        "increment_sequence",
        "increment_seq",
    )

    if func:
        return _call_best_effort(
            func,
            ham_int,
            sequence=sequence,
            increment=increment,
        )

    return _send_command_best_effort(
        ham_int,
        "SEQINCREMENT",
        sequenceObj=sequence,
        sequence=sequence,
        increment=increment,
    )


def ph_tip_pick_up_head_seq(ham_int, tip_seq: str, head_size: int = 96):
    head_size = int(head_size)

    if head_size not in {96, 384}:
        raise ValueError(f"Unsupported head size: {head_size}. Use 96 or 384.")

    func = _get_optional_pyhamilton_function(
        "tip_pick_up_head_seq",
        "tip_pick_up_384_seq" if head_size == 384 else "tip_pick_up_96_seq",
        "tip_pickup_head_seq",
        "tip_pickup_384_seq" if head_size == 384 else "tip_pickup_96_seq",
    )

    if func:
        return _call_best_effort(
            func,
            ham_int,
            tip_seq=tip_seq,
            tip384_seq=tip_seq,
            tip96_seq=tip_seq,
            tip_seq_name=tip_seq,
            sequence=tip_seq,
            head_size=head_size,
        )

    command_name = "PICKUP384" if head_size == 384 else "PICKUP96"

    return _send_command_best_effort(
        ham_int,
        command_name,
        tipSequence=tip_seq,
        channelVariable="1" * head_size,
        sequenceCounting=1,
    )


def ph_aspirate_head_seq(
    ham_int,
    plate,
    head_asp_seq: str,
    vols: float,
    liq_class: str,
    head_size: int = 96,
):
    head_size = int(head_size)
    liq_class = _clean_liquid_class(liq_class)

    if head_size not in {96, 384}:
        raise ValueError(f"Unsupported head size: {head_size}. Use 96 or 384.")

    print(f"USING {head_size}-HEAD ASPIRATE LIQUID CLASS: {liq_class}")

    func = _get_optional_pyhamilton_function(
        "aspirate_head_seq",
        "aspirate_384_seq" if head_size == 384 else "aspirate_96_seq",
        "aspirate384_seq" if head_size == 384 else "aspirate96_seq",
        "aspirate_384" if head_size == 384 else "aspirate_96",
        "aspirate384" if head_size == 384 else "aspirate96",
    )

    if func:
        return _call_best_effort(
            func,
            ham_int,
            plate,
            plate96=plate,
            plate384=plate,
            plate=plate,
            head_asp_seq=head_asp_seq,
            asp_seq=head_asp_seq,
            sequence=head_asp_seq,
            vols=vols,
            volume=vols,
            liq_class=liq_class,
            liquid_class=liq_class,
            liq_class2=liq_class,
            head_size=head_size,
        )

    command_name = "ASPIRATE384" if head_size == 384 else "ASPIRATE96"

    return _send_command_best_effort(
        ham_int,
        command_name,
        aspirateSequence=head_asp_seq,
        labwarePositions="",
        aspirateVolume=vols,
        volumes=vols,
        channelVariable="1" * head_size,
        liquidClass=liq_class,
        sequenceCounting=0,
        capacitiveLLD=2,
    )


def ph_dispense_head_seq(
    ham_int,
    plate,
    head_disp_seq: str,
    vols: float,
    liq_class: str,
    head_size: int = 96,
):
    head_size = int(head_size)
    liq_class = _clean_liquid_class(liq_class)

    if head_size not in {96, 384}:
        raise ValueError(f"Unsupported head size: {head_size}. Use 96 or 384.")

    print(f"USING {head_size}-HEAD DISPENSE LIQUID CLASS: {liq_class}")

    func = _get_optional_pyhamilton_function(
        "dispense_head_seq",
        "dispense_384_seq" if head_size == 384 else "dispense_96_seq2",
        "dispense_384_seq" if head_size == 384 else "dispense_96_seq",
        "dispense384_seq" if head_size == 384 else "dispense96_seq",
        "dispense_384" if head_size == 384 else "dispense_96",
        "dispense384" if head_size == 384 else "dispense96",
    )

    if func:
        return _call_best_effort(
            func,
            ham_int,
            plate,
            plate96=plate,
            plate384=plate,
            plate=plate,
            head_disp_seq=head_disp_seq,
            disp_seq=head_disp_seq,
            sequence=head_disp_seq,
            vols=vols,
            volume=vols,
            liq_class=liq_class,
            liquid_class=liq_class,
            liq_class2=liq_class,
            head_size=head_size,
        )

    command_name = "DISPENSE384" if head_size == 384 else "DISPENSE96"

    return _send_command_best_effort(
        ham_int,
        command_name,
        dispenseSequence=head_disp_seq,
        labwarePositions="",
        dispenseVolume=vols,
        volumes=vols,
        channelVariable="1" * head_size,
        liquidClass=liq_class,
        sequenceCounting=0,
    )


def ph_tip_eject_head(ham_int, head_size: int = 96):
    head_size = int(head_size)

    if head_size not in {96, 384}:
        raise ValueError(f"Unsupported head size: {head_size}. Use 96 or 384.")

    func = _get_optional_pyhamilton_function(
        "tip_eject_head",
        "tip_eject_384" if head_size == 384 else "tip_eject_96",
        "tip_eject384" if head_size == 384 else "tip_eject96",
        "tip_eject_384_seq" if head_size == 384 else "tip_eject_96_seq",
    )

    if func:
        return _call_best_effort(
            func,
            ham_int,
            head_size=head_size,
        )

    command_name = "EJECT384" if head_size == 384 else "EJECT96"

    return _send_command_best_effort(
        ham_int,
        command_name,
        labwarePositions="",
        channelVariable="1" * head_size,
        tipEjectToKnownPosition=2,
    )


def ph_tip_pick_up_96_seq(ham_int, tip96_seq: str):
    return ph_tip_pick_up_head_seq(
        ham_int,
        tip_seq=tip96_seq,
        head_size=96,
    )


def ph_aspirate_96_seq(ham_int, plate96, head_asp_seq: str, vols: float, liq_class: str):
    return ph_aspirate_head_seq(
        ham_int,
        plate=plate96,
        head_asp_seq=head_asp_seq,
        vols=vols,
        liq_class=liq_class,
        head_size=96,
    )


def ph_dispense_96_seq2(ham_int, plate96, head_disp_seq: str, vols: float, liq_class: str):
    return ph_dispense_head_seq(
        ham_int,
        plate=plate96,
        head_disp_seq=head_disp_seq,
        vols=vols,
        liq_class=liq_class,
        head_size=96,
    )


def ph_tip_eject_96(ham_int):
    return ph_tip_eject_head(
        ham_int,
        head_size=96,
    )


def ph_tip_pick_up_384_seq(ham_int, tip384_seq: str):
    return ph_tip_pick_up_head_seq(
        ham_int,
        tip_seq=tip384_seq,
        head_size=384,
    )


def ph_aspirate_384_seq(ham_int, plate384, head_asp_seq: str, vols: float, liq_class: str):
    return ph_aspirate_head_seq(
        ham_int,
        plate=plate384,
        head_asp_seq=head_asp_seq,
        vols=vols,
        liq_class=liq_class,
        head_size=384,
    )


def ph_dispense_384_seq(ham_int, plate384, head_disp_seq: str, vols: float, liq_class: str):
    return ph_dispense_head_seq(
        ham_int,
        plate=plate384,
        head_disp_seq=head_disp_seq,
        vols=vols,
        liq_class=liq_class,
        head_size=384,
    )


def ph_tip_eject_384(ham_int):
    return ph_tip_eject_head(
        ham_int,
        head_size=384,
    )


@dataclass
class SequenceStepConfig:
    sequence: str
    role: str
    liquid_class: str
    transfer_type: str
    replicate: str
    volume_ul: float
    marker: str
    channel_pattern: str = ""
    hardware_mode: str = "channels"
    selected_index: int = 0
    sequence_count: int = 0
    manual: bool = False
    set_order: int = 1
    control: bool = False


@dataclass
class TransferPair:
    pair_id: str
    tip_sequence: str = ""
    source_sequence: str = ""
    destination_sequence: str = ""
    source_position: str = ""
    destination_position: str = ""
    source_sample_number: Optional[int] = None
    destination_sample_number: Optional[int] = None
    replicate_index: int = 1
    volume_ul: float = 0.0
    liquid_class: str = ""
    transfer_type: str = ""
    marker_label: str = ""


@dataclass
class ValidationIssue:
    severity: str
    category: str
    message: str
    evidence: Dict[str, Any] = field(default_factory=dict)


@dataclass
class DevPalLiteOutput:
    lay_file: str
    layout_file: str
    trace_file: str
    selected_sequences: List[Dict[str, Any]]
    source_positions: List[str]
    destination_positions: List[str]
    transfer_pairs: List[Dict[str, Any]]
    execution_plan: List[Dict[str, Any]]
    marker_matches: Dict[str, List[Dict[str, Any]]]
    layout_findings: List[Dict[str, Any]]
    trace_findings: List[Dict[str, Any]]
    validation_findings: List[Dict[str, Any]]
    issues: List[Dict[str, Any]]
    limitations: List[str]
    generated_review_script: str
    report_path: str


class AutoRunHamiltonInterface(phi.HamiltonInterface if phi else object):
    def start(self):
        if phi is None:
            raise RuntimeError(f"pyhamilton is not available. Import error: {PYHAMILTON_IMPORT_ERROR}")

        if self.active:
            return

        self.log("starting Hamilton interface")

        if self.simulate:
            self.server_thread = phi.HamiltonInterface.HamiltonServerThread(
                self.address,
                self.port,
            )
            self.server_thread.start()
            time.sleep(1)

            subprocess.Popen([
                OEM_RUN_EXE_PATH,
                OEM_HSL_PATH,
                "-r",
                "-t",
            ])

            self.log("started OEM application for simulation with auto-run")
        else:
            self.oem_process = Process(target=phi.run_hamilton_process, args=())
            self.oem_process.start()
            self.log("started OEM process")

            self.server_thread = phi.HamiltonInterface.HamiltonServerThread(
                self.address,
                self.port,
            )
            self.server_thread.start()

        self.log("started server thread")
        self.active = True


class SearchableCombobox(ttk.Combobox):
    def __init__(self, master=None, completevalues=None, **kwargs):
        self.completevalues = sorted(completevalues or [], key=str.lower)
        super().__init__(master, values=self.completevalues, **kwargs)
        self.bind("<KeyRelease>", self._filter_values)

    def _filter_values(self, event):
        typed = self.get().strip().lower()
        if not typed:
            self["values"] = self.completevalues
            return

        filtered = [value for value in self.completevalues if typed in value.lower()]
        self["values"] = filtered

        if filtered:
            try:
                self.event_generate("<Down>")
            except Exception:
                pass


def safe_float(value: Any, default: float = 0.0) -> float:
    try:
        return float(value)
    except Exception:
        return default


def normalize_hardware_mode(value: str) -> str:
    text = str(value or "").strip().lower().replace(" ", "").replace("_", "-")

    if text in {"head384", "384head", "384", "head-384", "mph384", "384mph"}:
        return "head384"

    if text in {"head96", "96head", "head", "96", "head-96", "mph96", "96mph"}:
        return "head96"

    return "channels"


def get_head_size_from_mode(hardware_mode: str) -> int:
    mode = normalize_hardware_mode(hardware_mode)
    if mode == "head384":
        return 384
    if mode == "head96":
        return 96
    return 0


def normalize_assigned_channel_pattern(channel_pattern: str, hardware_mode: str = "channels") -> str:
    mode = normalize_hardware_mode(hardware_mode)
    pattern = str(channel_pattern or "").strip()

    if mode == "head384":
        return "1" * 384

    if mode == "head96":
        return "1" * 96

    if not re.match(r"^[01]+$", pattern):
        raise ValueError(f"Invalid channel pattern: {pattern}")

    if len(pattern) not in VALID_CHANNEL_PATTERN_LENGTHS:
        raise ValueError(
            f"Invalid channel pattern length {len(pattern)}. "
            "Use 8, 12, 14, or 16 digits for channel mode."
        )

    if "1" not in pattern:
        raise ValueError("Channel pattern must contain at least one active channel.")

    return pattern


def excel_col_name(index_zero_based: int) -> str:
    name = ""
    index = index_zero_based + 1

    while index:
        index, remainder = divmod(index - 1, 26)
        name = chr(65 + remainder) + name

    return name


def normalize_position_text(value: Any) -> str:
    text = str(value).strip().upper()
    text = text.replace(" ", "")
    text = text.replace(":", "")
    return text

def infer_plate_format_from_dimensions(row_count: int, col_count: int) -> str:
    if row_count <= 4 and col_count <= 6:
        return "24-well"
    if row_count <= 8 and col_count <= 12:
        return "96-well"
    if row_count <= 16 and col_count <= 24:
        return "384-well"
    return f"unknown-{row_count}x{col_count}"


def normalize_lay_text(text: str) -> str:
    text = str(text or "")
    text = text.replace("\x00", " ")
    text = re.sub(r"[\x01-\x1f]+", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def clean_lay_value(value: Any) -> str:
    value = clean_control_chars(value)
    value = value.strip(" ;:=[]{}'\"")
    value = re.sub(r"\s+", " ", value)
    return value.strip()


def infer_plate_format_from_position_count(count: int) -> str:
    return "unknown"


def is_carrier_labware(labware_id: str, labware_file: str, properties: Dict[str, Any]) -> bool:
    text = " ".join([
        str(labware_id or ""),
        str(labware_file or ""),
        os.path.basename(str(labware_file or "")),
        " ".join(str(key) for key in properties.keys()),
        " ".join(str(value) for value in properties.values()),
    ]).upper()

    return (
        str(labware_file or "").lower().endswith(".tml")
        or "CARRIER" in text
        or "MLSTARCAR" in text
        or "TIP_CAR" in text
        or "PLT_CAR" in text
        or "SMP_CAR" in text
        or "TAB" in text
    )


def infer_labware_format_from_labware_record(
    labware_id: str,
    labware_file: str = "",
    properties: Optional[Dict[str, Any]] = None,
) -> str:
    properties = properties or {}

    text = " ".join([
        str(labware_id or ""),
        str(labware_file or ""),
        os.path.basename(str(labware_file or "")),
        " ".join(str(key) for key in properties.keys()),
        " ".join(str(value) for value in properties.values()),
    ]).upper()

    if is_carrier_labware(labware_id, labware_file, properties):
        if "TIP_CAR" in text or "TIPCAR" in text:
            return "tip carrier"
        if "PLT_CAR" in text or "PLTCAR" in text:
            return "plate carrier"
        if "SMP_CAR" in text or "SAMPLE_CAR" in text:
            return "sample carrier"
        return "carrier"

    if "SLIMTIP300" in text or "SLIMTIP" in text or "300UL" in text or "300S" in text:
        return "300 uL slim/filter tip rack"

    if "HTF" in text or "HVT" in text or "HIGHVOLUME" in text or "HIGH_VOLUME" in text or "1000UL" in text:
        return "1000 uL high-volume/filter tip rack"

    if "50UL" in text or "LOWVOLUME" in text or "LOW_VOLUME" in text or "LVT" in text:
        return "50 uL low-volume tip rack"

    if "WASTE" in text:
        return "waste"

    if "HONEYCOMB" in text or "48_TUBE" in text or "48TUBE" in text or (re.search(r"\b48\b", text) and "TUBE" in text):
        return "48-position tube/sample rack"

    if "MAGMAX24" in text or ("24" in text and ("PLATE" in text or "WELL" in text)):
        return "24-well plate"

    if "384" in text:
        return "384-well plate"

    if "96" in text and any(token in text for token in ("PLATE", "NUNC", "RGT", "DWP", "MTP", "WELL", "96L")):
        return "96-well plate"

    if "TIP" in text:
        return "tip rack"

    if "PLATE" in text:
        return "plate/unknown format"

    if "RACK" in text or "TUBE" in text:
        return "rack/unknown capacity"

    return "unknown"


def infer_plate_format_from_labware_name(name: str) -> str:
    return infer_labware_format_from_labware_record(str(name or ""), "", {})


def infer_labware_format_from_name(labware_name: str) -> str:
    return infer_labware_format_from_labware_record(str(labware_name or ""), "", {})


def parse_lay_labware_catalog(lay_path: str) -> Dict[str, Dict[str, Any]]:
    text = normalize_lay_text(read_text_file_best_effort(lay_path))
    catalog = {}

    if not text:
        return catalog

    labware_nums = sorted(
        {int(value) for value in re.findall(r"Labware\.(\d+)\.", text, re.IGNORECASE)}
    )

    for labware_num in labware_nums:
        prefix = rf"Labware\.{labware_num}\."

        def field_value(field: str) -> str:
            pattern = re.compile(
                prefix + re.escape(field) + r"\s*(.*?)(?=\s+Labware\.\d+\.|\s+Layer\.\d+\.|\s+Seq\.\d+\.|\s+Labware\.Cnt|\s+Layer\.Cnt|\s+Seq\.Cnt|\s+PhoenixVersion|\s+UseGlobalTpl|$)",
                re.IGNORECASE,
            )
            match = pattern.search(text)
            return clean_lay_value(match.group(1)) if match else ""

        labware_id = field_value("Id")
        labware_file = field_value("File")
        template = field_value("Template")
        site_id = field_value("SiteId")

        properties = {}
        property_matches = re.findall(
            prefix + r"LwProperties\.(\d+)\.Property\s*(.*?)\s+"
            + prefix + r"LwProperties\.\1\.Value\s*(.*?)(?=\s+Labware\.\d+\.|\s+Layer\.|\s+Seq\.|\s+Labware\.Cnt|$)",
            text,
            re.IGNORECASE,
        )

        for _, prop_name, prop_value in property_matches:
            prop_name = clean_lay_value(prop_name)
            prop_value = clean_lay_value(prop_value)
            if prop_name:
                properties[prop_name] = prop_value

        if not labware_id:
            continue

        labware_format = infer_labware_format_from_labware_record(
            labware_id,
            labware_file,
            properties,
        )

        record = {
            "labware_number": labware_num,
            "labware_id": labware_id,
            "labware_file": labware_file,
            "labware_file_name": os.path.basename(labware_file.replace("\\", os.sep)) if labware_file else "",
            "template": template,
            "site_id": site_id,
            "properties": properties,
            "is_carrier": is_carrier_labware(labware_id, labware_file, properties),
            "labware_format": labware_format,
            "format_basis": "Labware.N.Id/File/Template/Properties",
        }

        catalog[labware_id] = record
        catalog[labware_id.lower()] = record

    return catalog


def parse_lay_layer_sequence_map(lay_path: str) -> Dict[str, str]:
    text = normalize_lay_text(read_text_file_best_effort(lay_path))
    seq_to_labware = {}

    matches = re.findall(
        r"Layer\.\d+\.(\d+)\.LabwareName\s*(.*?)\s+Layer\.\d+\.\1\.SeqName\s*(.*?)(?=\s+Layer\.|\s+Labware\.|\s+Seq\.|$)",
        text,
        re.IGNORECASE,
    )

    for _, labware_name, seq_name in matches:
        labware_name = clean_lay_value(labware_name)
        seq_name = clean_lay_value(seq_name)

        if seq_name and labware_name:
            seq_to_labware[seq_name] = labware_name
            seq_to_labware[seq_name.lower()] = labware_name

    return seq_to_labware


def parse_lay_file(lay_path: str) -> List[str]:
    text = normalize_lay_text(read_text_file_best_effort(lay_path))

    if not text:
        return []

    found = []
    seen = set()

    seq_ids = sorted(
        {int(value) for value in re.findall(r"Seq\.(\d+)\.Name", text, re.IGNORECASE)}
    )

    for seq_id in seq_ids:
        match = re.search(
            rf"Seq\.{seq_id}\.Name\s*(.*?)(?=\s+Seq\.{seq_id}\.ReadOnly|\s+Seq\.\d+\.|\s+Seq\.Cnt|\s+Layer\.|\s+Labware\.|$)",
            text,
            re.IGNORECASE,
        )

        if not match:
            continue

        name = clean_lay_value(match.group(1))

        if not name or name.lower().startswith(("readonly", "cnt", "item")):
            continue

        if name not in seen:
            found.append(name)
            seen.add(name)

    return found


def parse_lay_sequence_counts(lay_path: str) -> Dict[str, int]:
    text = normalize_lay_text(read_text_file_best_effort(lay_path))

    if not text:
        return {}

    counts_by_num = {}
    names_by_num = {}

    for seq_num, cnt in re.findall(r"Seq\.(\d+)\.Cnt[^\d]*(\d+)", text, re.IGNORECASE):
        counts_by_num[int(seq_num)] = int(cnt)

    seq_ids = sorted(
        {int(value) for value in re.findall(r"Seq\.(\d+)\.Name", text, re.IGNORECASE)}
    )

    for seq_id in seq_ids:
        match = re.search(
            rf"Seq\.{seq_id}\.Name\s*(.*?)(?=\s+Seq\.{seq_id}\.ReadOnly|\s+Seq\.\d+\.|\s+Seq\.Cnt|\s+Layer\.|\s+Labware\.|$)",
            text,
            re.IGNORECASE,
        )

        if match:
            name = clean_lay_value(match.group(1))
            if name:
                names_by_num[seq_id] = name

    return {
        name: counts_by_num.get(seq_num, 0)
        for seq_num, name in names_by_num.items()
    }


def parse_lay_sequence_metadata(lay_path: str) -> Dict[str, Dict[str, Any]]:
    text = normalize_lay_text(read_text_file_best_effort(lay_path))
    metadata = {}

    if not text:
        return metadata

    labware_catalog = parse_lay_labware_catalog(lay_path)
    layer_sequence_map = parse_lay_layer_sequence_map(lay_path)

    seq_ids = sorted(set(re.findall(r"Seq\.(\d+)\.", text)), key=lambda value: int(value))

    for seq_id in seq_ids:
        name_match = re.search(
            rf"Seq\.{seq_id}\.Name\s*(.*?)(?=\s+Seq\.{seq_id}\.ReadOnly|\s+Seq\.\d+\.|\s+Seq\.Cnt|\s+Layer\.|\s+Labware\.|$)",
            text,
            re.IGNORECASE,
        )

        if not name_match:
            continue

        seq_name = clean_lay_value(name_match.group(1))

        if not seq_name:
            continue

        cnt_match = re.search(rf"Seq\.{seq_id}\.Cnt[^\d]*(\d+)", text, re.IGNORECASE)
        selected_position_count = int(cnt_match.group(1)) if cnt_match else 0

        item_pattern = re.compile(
            rf"Seq\.{seq_id}\.Item\.(\d+)\.ObjId\s*(.*?)\s+Seq\.{seq_id}\.Item\.\1\.PosId\s*([A-Za-z]?\d{{1,3}})",
            re.IGNORECASE,
        )

        objids = []
        positions = []

        for _, objid, posid in item_pattern.findall(text):
            objid = clean_lay_value(objid)
            posid = normalize_position_text(posid)

            if objid:
                objids.append(objid)
            if posid:
                positions.append(posid)

        layer_labware_name = (
            layer_sequence_map.get(seq_name)
            or layer_sequence_map.get(seq_name.lower())
            or ""
        )

        primary_objid = max(set(objids), key=objids.count) if objids else layer_labware_name

        labware_record = (
            labware_catalog.get(primary_objid)
            or labware_catalog.get(str(primary_objid).lower())
            or labware_catalog.get(layer_labware_name)
            or labware_catalog.get(str(layer_labware_name).lower())
            or {}
        )

        labware_name = (
            labware_record.get("labware_id")
            or primary_objid
            or layer_labware_name
            or seq_name
        )

        labware_file = labware_record.get("labware_file", "")
        labware_format = labware_record.get("labware_format", "unknown")

        if labware_format == "unknown":
            labware_format = infer_labware_format_from_labware_record(
                labware_name,
                labware_file,
                labware_record.get("properties", {}),
            )

        metadata[seq_name] = {
            "sequence_id": seq_id,
            "sequence_name": seq_name,
            "selected_position_count": selected_position_count,
            "position_count": selected_position_count,
            "positions": positions,
            "labware_name": labware_name,
            "primary_objid": primary_objid,
            "layer_labware_name": layer_labware_name,
            "all_detected_labware_names": sorted(
                {value for value in objids + [primary_objid, layer_labware_name, labware_name] if value},
                key=str.lower,
            ),
            "labware_file": labware_file,
            "labware_file_name": labware_record.get("labware_file_name", ""),
            "labware_template": labware_record.get("template", ""),
            "labware_site_id": labware_record.get("site_id", ""),
            "labware_properties": labware_record.get("properties", {}),
            "is_carrier": labware_record.get("is_carrier", False),
            "labware_format": labware_format,
            "labware_format_basis": "Seq.Item.ObjId joined to Labware.N.Id/File/Template/Properties and Layer map",
        }

    return metadata


def get_sequence_labware_metadata(
    sequence_name: str,
    lay_metadata: Dict[str, Dict[str, Any]],
    sequence_count: int = 0,
) -> Dict[str, Any]:
    target = str(sequence_name or "").strip()
    target_lower = target.lower()

    if target in lay_metadata:
        return lay_metadata[target]

    for name, metadata in lay_metadata.items():
        if str(name).strip().lower() == target_lower:
            return metadata

    inferred_format = infer_labware_format_from_labware_record(target, "", {})

    return {
        "sequence_name": target,
        "labware_name": target if target else "unknown",
        "primary_objid": target,
        "layer_labware_name": "",
        "all_detected_labware_names": [target] if target else [],
        "labware_file": "",
        "labware_file_name": "",
        "labware_template": "",
        "labware_site_id": "",
        "labware_properties": {},
        "is_carrier": False,
        "labware_format": inferred_format,
        "labware_format_basis": "sequence-name fallback only; method labware record not found",
        "selected_position_count": int(sequence_count or 0),
        "position_count": int(sequence_count or 0),
        "positions": [],
    }


def normalize_tip_type(value: str) -> str:
    text = str(value or "").lower().replace("_", "").replace("-", "").replace(" ", "")

    if text in {"1000ul", "1000", "hv", "hvt", "highvolume", "highvolumefilter"}:
        return "1000ul"

    if text in {"300ul", "300", "300s", "std", "slim", "slimfilter"}:
        return "300ul"

    if text in {"50ul", "50", "lv", "lvt", "lowvolume"}:
        return "50ul"

    return ""


def infer_tip_type_from_text(text: str) -> str:
    text = str(text or "").upper()

    if "SLIMTIP300" in text or "SLIM" in text or "300UL" in text or "300S" in text or re.search(r"\b300\b", text):
        return "300ul"

    if "HTF" in text or "HVT" in text or "HIGHVOLUME" in text or "HIGH_VOLUME" in text or "1000UL" in text or re.search(r"\b1000\b", text):
        return "1000ul"

    if "50UL" in text or "LOWVOLUME" in text or "LOW_VOLUME" in text or "LVT" in text or re.search(r"\b50\b", text):
        return "50ul"

    return ""


def infer_tip_type_for_set(
    set_order: int,
    configs: List[SequenceStepConfig],
    lay_metadata: Dict[str, Dict[str, Any]],
) -> Tuple[str, str]:
    tip_configs = [
        config for config in configs
        if int(config.set_order) == int(set_order)
        and config.role == "tip pick up"
    ]

    searched = []

    for config in tip_configs:
        meta = get_sequence_labware_metadata(
            config.sequence,
            lay_metadata,
            getattr(config, "sequence_count", 0),
        )

        values = [
            config.sequence,
            meta.get("sequence_name", ""),
            meta.get("labware_name", ""),
            meta.get("primary_objid", ""),
            meta.get("layer_labware_name", ""),
            meta.get("labware_format", ""),
            meta.get("labware_file", ""),
            meta.get("labware_file_name", ""),
            meta.get("labware_template", ""),
            meta.get("labware_site_id", ""),
            " ".join(meta.get("all_detected_labware_names", []) or []),
        ]

        for key, value in (meta.get("labware_properties", {}) or {}).items():
            values.append(str(key))
            values.append(str(value))

        for value in values:
            text = str(value or "").strip()
            if not text:
                continue

            searched.append(text)
            tip_type = normalize_tip_type(infer_tip_type_from_text(text))

            if tip_type in {"1000ul", "300ul", "50ul"}:
                return tip_type, text

    return "", " | ".join(dict.fromkeys(searched)) if searched else "no tip pick up sequence selected"


def tip_volume_allowed(tip_type: str, hardware_mode: str) -> Optional[Tuple[float, float]]:
    tip_type = normalize_tip_type(tip_type)
    mode = normalize_hardware_mode(hardware_mode)

    is_head = mode in {"head96", "head384"}

    if is_head:
        ranges = {
            "1000ul": (100.0, 1000.0),
            "300ul": (20.0, 300.0),
            "50ul": (12.5, 49.0),
        }
    else:
        ranges = {
            "1000ul": (100.0, 1000.0),
            "300ul": (10.0, 300.0),
            "50ul": (10.0, 50.0),
        }

    return ranges.get(tip_type)


def validate_tip_volume_rules(
    configs: List[SequenceStepConfig],
    lay_metadata: Dict[str, Dict[str, Any]],
) -> List[ValidationIssue]:
    issues = []
    grouped = defaultdict(list)

    for config in configs:
        grouped[int(getattr(config, "set_order", 1) or 1)].append(config)

    for set_order, items in sorted(grouped.items()):
        liquid_steps = [item for item in items if item.role in ("aspirate", "dispense")]

        if not liquid_steps:
            continue

        tip_type, evidence = infer_tip_type_for_set(set_order, configs, lay_metadata)

        if not tip_type:
            issues.append(ValidationIssue(
                "major",
                "tip_type",
                f"Could not identify tip type for Set {set_order}. Checked method/deck metadata: {evidence}",
                {"set_order": set_order, "checked_metadata": evidence},
            ))
            continue

        hardware_mode = normalize_hardware_mode(liquid_steps[0].hardware_mode)
        allowed_range = tip_volume_allowed(tip_type, hardware_mode)

        if not allowed_range:
            continue

        min_ul, max_ul = allowed_range

        for step in liquid_steps:
            volume = float(step.volume_ul)

            if volume < min_ul or volume > max_ul:
                issues.append(ValidationIssue(
                    "critical",
                    "tip_volume_range",
                    (
                        f"Set {set_order} volume {volume} uL is outside allowed range "
                        f"{min_ul}-{max_ul} uL for {hardware_mode} {tip_type} tips."
                    ),
                    {
                        "set_order": set_order,
                        "sequence": step.sequence,
                        "tip_type": tip_type,
                        "tip_evidence": evidence,
                        "hardware_mode": hardware_mode,
                        "volume_ul": volume,
                        "allowed_min_ul": min_ul,
                        "allowed_max_ul": max_ul,
                    },
                ))

    return issues


def unique_issue_messages(issues: List[ValidationIssue]) -> List[str]:
    seen = set()
    messages = []

    for issue in issues:
        message = str(issue.message or "").strip()
        if message and message not in seen:
            seen.add(message)
            messages.append(message)

    return messages


def validate_step_configs(configs: List[SequenceStepConfig]) -> List[ValidationIssue]:
    issues = []
    roles = {config.role for config in configs}

    for required_role in ROLE_OPTIONS:
        if required_role not in roles:
            issues.append(ValidationIssue("critical", "configuration", f"Missing required step role: {required_role}"))

    for config in configs:
        if config.role not in ROLE_OPTIONS:
            issues.append(ValidationIssue("critical", "configuration", f"Invalid step role for {config.sequence}: {config.role}"))

        if config.role in ("aspirate", "dispense") and not config.liquid_class:
            issues.append(ValidationIssue("critical", "configuration", f"Liquid class is required for {config.role}: {config.sequence}"))

        if config.volume_ul <= 0 or config.volume_ul > 1000:
            issues.append(ValidationIssue("critical", "configuration", f"Volume must be between 1 and 1000 uL: {config.sequence}"))

        if config.transfer_type not in TRANSFER_OPTIONS:
            issues.append(ValidationIssue("critical", "configuration", f"Invalid transfer type: {config.transfer_type}"))

        if config.replicate not in REPLICATE_OPTIONS:
            issues.append(ValidationIssue("critical", "configuration", f"Invalid replicate value: {config.replicate}"))

        if config.transfer_type == "samples" and not config.marker:
            issues.append(ValidationIssue("critical", "configuration", f"Marker is required for samples: {config.sequence}"))

    return issues


def get_sequence_batch_size(config: SequenceStepConfig) -> int:
    mode = normalize_hardware_mode(config.hardware_mode)

    if mode == "head384":
        return 384

    if mode == "head96":
        return 96

    pattern = normalize_assigned_channel_pattern(config.channel_pattern, "channels")
    return max(pattern.count("1"), 1)


def get_sequence_iterations(config: SequenceStepConfig) -> int:
    if bool(config.manual):
        return 1

    position_count = int(config.sequence_count or 0)
    batch_size = get_sequence_batch_size(config)

    if position_count <= 0:
        return 1

    return max((position_count + batch_size - 1) // batch_size, 1)


def any_control_sequence_selected(configs: List[SequenceStepConfig]) -> bool:
    return any(bool(getattr(config, "control", False)) for config in configs)


def get_required_transfer_iterations(configs: List[SequenceStepConfig]) -> int:
    control_configs = [
        config for config in configs
        if bool(getattr(config, "control", False))
        and not bool(getattr(config, "manual", False))
        and config.role in ("aspirate", "dispense")
    ]

    if control_configs:
        return max(get_sequence_iterations(config) for config in control_configs)

    nonmanual_liquid_configs = [
        config for config in configs
        if config.role in ("aspirate", "dispense")
        and not bool(getattr(config, "manual", False))
    ]

    if nonmanual_liquid_configs:
        return max(get_sequence_iterations(config) for config in nonmanual_liquid_configs)

    return 1


def build_execution_plan(
    configs: List[SequenceStepConfig],
    lay_metadata: Optional[Dict[str, Dict[str, Any]]] = None,
) -> List[Dict[str, Any]]:
    lay_metadata = lay_metadata or {}
    plan = []
    order = 1

    grouped = defaultdict(list)
    for config in configs:
        grouped[int(getattr(config, "set_order", 1) or 1)].append(config)

    for set_order in sorted(grouped):
        set_configs = sorted(grouped[set_order], key=lambda c: c.selected_index)

        replicate = set_configs[0].replicate if set_configs else "single"
        replicate_count = REPLICATE_MAP.get(str(replicate).lower(), 1)

        tip_steps = [c for c in set_configs if c.role == "tip pick up"]
        asp_steps = [c for c in set_configs if c.role == "aspirate"]
        dsp_steps = [c for c in set_configs if c.role == "dispense"]

        max_len = max(len(tip_steps), len(asp_steps), len(dsp_steps), 1)
        transfer_iterations = get_required_transfer_iterations(set_configs)

        for rep_index in range(1, replicate_count + 1):
            for transfer_iteration in range(1, transfer_iterations + 1):
                for idx in range(max_len):
                    triplet = [
                        tip_steps[idx] if idx < len(tip_steps) else tip_steps[0] if tip_steps else None,
                        asp_steps[idx] if idx < len(asp_steps) else asp_steps[0] if asp_steps else None,
                        dsp_steps[idx] if idx < len(dsp_steps) else dsp_steps[0] if dsp_steps else None,
                    ]

                    for config in triplet:
                        if config is None:
                            continue

                        batch_size = get_sequence_batch_size(config)
                        action = "tip_pick" if config.role == "tip pick up" else config.role

                        should_increment = (
                            transfer_iteration > 1
                            and not bool(getattr(config, "manual", False))
                            and config.role in ("aspirate", "dispense")
                        )

                        seq_meta = get_sequence_labware_metadata(
                            config.sequence,
                            lay_metadata,
                            getattr(config, "sequence_count", 0),
                        )

                        plan.append({
                            "order": order,
                            "set_order": set_order,
                            "action": action,
                            "sequence": config.sequence,
                            "hardware_mode": config.hardware_mode,
                            "labware_name": seq_meta.get("labware_name", ""),
                            "labware_format": seq_meta.get("labware_format", ""),
                            "labware_positions": seq_meta.get("positions", []),
                            "replicate_index": rep_index,
                            "transfer_iteration": transfer_iteration,
                            "manual": bool(getattr(config, "manual", False)),
                            "control": bool(getattr(config, "control", False)),
                            "autoincrement": should_increment,
                            "increment": batch_size if should_increment else 0,
                            "channel_pattern": config.channel_pattern,
                            "volume_ul": config.volume_ul if config.role in ("aspirate", "dispense") else "",
                            "liquid_class": config.liquid_class if config.role in ("aspirate", "dispense") else "",
                            "transfer_type": config.transfer_type,
                            "marker": config.marker,
                            "sequence_count": config.sequence_count,
                            "batch_size": batch_size,
                        })
                        order += 1

                    if tip_steps:
                        tip = tip_steps[idx] if idx < len(tip_steps) else tip_steps[0]

                        plan.append({
                            "order": order,
                            "set_order": set_order,
                            "action": "tip_eject",
                            "sequence": "Waste",
                            "hardware_mode": tip.hardware_mode,
                            "labware_name": "Waste",
                            "labware_format": "waste",
                            "labware_positions": [],
                            "replicate_index": rep_index,
                            "transfer_iteration": transfer_iteration,
                            "manual": False,
                            "control": False,
                            "autoincrement": False,
                            "increment": 0,
                            "channel_pattern": tip.channel_pattern,
                            "volume_ul": "",
                            "liquid_class": "",
                            "transfer_type": tip.transfer_type,
                            "marker": tip.marker,
                            "sequence_count": tip.sequence_count,
                            "batch_size": get_sequence_batch_size(tip),
                        })
                        order += 1

    return plan


def generate_deterministic_transfer_pairs(configs: List[SequenceStepConfig], layout_matches: Dict[str, List[Dict[str, Any]]]) -> List[TransferPair]:
    replicate = configs[0].replicate if configs else "single"
    replicate_count = REPLICATE_MAP.get(str(replicate).lower(), 1)

    tip_steps = [c for c in configs if c.role == "tip pick up"]
    asp_steps = [c for c in configs if c.role == "aspirate"]
    dsp_steps = [c for c in configs if c.role == "dispense"]

    if not dsp_steps:
        return []

    marker = configs[0].marker.upper().strip() if configs and configs[0].marker else ""
    pairs = []
    pair_counter = 1

    max_dispense_positions = max(int(dsp.sequence_count or 0) for dsp in dsp_steps)
    if max_dispense_positions <= 0:
        max_dispense_positions = 1

    for rep_index in range(1, replicate_count + 1):
        for dest_index in range(1, max_dispense_positions + 1):
            asp = asp_steps[(dest_index - 1) % len(asp_steps)] if asp_steps else None
            dsp = dsp_steps[(dest_index - 1) % len(dsp_steps)] if dsp_steps else None
            tip = tip_steps[(dest_index - 1) % len(tip_steps)] if tip_steps else None

            source_num = dest_index
            source_position = str(source_num)
            marker_label = f"{marker}{source_num}" if marker else ""

            dest_position = ""
            if marker_label and marker_label in layout_matches:
                entries = layout_matches.get(marker_label, [])
                if entries:
                    dest_entry_index = min(rep_index - 1, len(entries) - 1)
                    dest_position = entries[dest_entry_index].get("dest_value", "")

            pairs.append(TransferPair(
                pair_id=f"pair_{pair_counter:05d}",
                tip_sequence=tip.sequence if tip else "",
                source_sequence=asp.sequence if asp else "",
                destination_sequence=dsp.sequence if dsp else "",
                source_position=source_position,
                destination_position=dest_position,
                source_sample_number=source_num,
                destination_sample_number=extract_position_number(dest_position),
                replicate_index=rep_index,
                volume_ul=asp.volume_ul if asp else dsp.volume_ul if dsp else 0.0,
                liquid_class=asp.liquid_class if asp else dsp.liquid_class if dsp else "",
                transfer_type=asp.transfer_type if asp else dsp.transfer_type if dsp else "",
                marker_label=marker_label,
            ))

            pair_counter += 1

    return pairs


def parse_dataframe_plate_layout(df, marker_pattern, plate_layout):
    if df is None or df.empty:
        return

    header_row_idx = 0
    row_label_col_idx = 0

    for row_idx in range(1, df.shape[0]):
        plate_row_raw = df.iat[row_idx, row_label_col_idx]

        if pd.isna(plate_row_raw):
            continue

        plate_row = str(plate_row_raw).strip().upper()

        if not re.match(r"^[A-P]$", plate_row):
            continue

        for col_idx in range(1, df.shape[1]):
            plate_col_raw = df.iat[header_row_idx, col_idx]

            if pd.isna(plate_col_raw):
                continue

            plate_col_text = str(plate_col_raw).strip()

            plate_col_match = re.search(r"\d+", plate_col_text)
            if not plate_col_match:
                continue

            plate_col = int(plate_col_match.group(0))

            cell_value = df.iat[row_idx, col_idx]

            if pd.isna(cell_value):
                continue

            cell_text = str(cell_value).strip().upper()

            if not cell_text:
                continue

            matches = marker_pattern.findall(cell_text)

            if not matches:
                continue

            dest_position = f"{plate_row}{plate_col}"

            for found_marker in matches:
                plate_layout[found_marker.upper()].append({
                    "dest_value": dest_position,
                    "cell_value": cell_text,
                    "plate_row": plate_row,
                    "plate_column": plate_col,
                    "row_index": row_idx + 1,
                    "column_index": col_idx + 1,
                    "excel_position": f"{excel_col_name(col_idx)}{row_idx + 1}",
                    "column_name": excel_col_name(col_idx),
                })


def parse_plate_layout_file(layout_path: str, marker: str) -> Dict[str, List[Dict[str, Any]]]:
    if pd is None:
        raise RuntimeError("pandas is required for layout parsing.")

    marker = marker.strip().upper()
    if not marker:
        return {}

    marker_pattern = re.compile(rf"\b{re.escape(marker)}\d+\b", re.IGNORECASE)
    plate_layout = defaultdict(list)

    ext = os.path.splitext(layout_path)[1].lower()

    if ext == ".csv":
        df = pd.read_csv(layout_path, header=None, dtype=str)
        parse_dataframe_plate_layout(df, marker_pattern, plate_layout)

    elif ext in (".xls", ".xlsx"):
        sheets = pd.read_excel(layout_path, sheet_name=None, header=None, dtype=str)
        for _, df in sheets.items():
            parse_dataframe_plate_layout(df, marker_pattern, plate_layout)


    elif ext == ".docx":

        if Document is None:
            raise RuntimeError("python-docx is required for DOCX layout parsing.")

        doc = Document(layout_path)

        for table in doc.tables:

            if not table.rows:
                continue

            header_cells = table.rows[0].cells

            for row_idx in range(1, len(table.rows)):

                row = table.rows[row_idx]

                plate_row = str(row.cells[0].text).strip().upper()

                if not re.match(r"^[A-P]$", plate_row):
                    continue

                for col_idx in range(1, len(row.cells)):

                    if col_idx >= len(header_cells):
                        continue

                    plate_col_text = str(header_cells[col_idx].text).strip()

                    plate_col_match = re.search(r"\d+", plate_col_text)

                    if not plate_col_match:
                        continue

                    plate_col = int(plate_col_match.group(0))

                    cell_text = str(row.cells[col_idx].text).strip().upper()

                    matches = marker_pattern.findall(cell_text)

                    if not matches:
                        continue

                    dest_position = f"{plate_row}{plate_col}"

                    for found_marker in matches:
                        plate_layout[found_marker.upper()].append({

                            "dest_value": dest_position,

                            "cell_value": cell_text,

                            "plate_row": plate_row,

                            "plate_column": plate_col,

                            "row_index": row_idx + 1,

                            "column_index": col_idx + 1,

                            "excel_position": f"{excel_col_name(col_idx)}{row_idx + 1}",

                            "column_name": excel_col_name(col_idx),

                        })
    else:
        raise ValueError("Unsupported plate layout file type. Use CSV, XLS, XLSX, or DOCX.")

    return dict(plate_layout)


def build_labware_validation_summary(
    layout_matches: Dict[str, List[Dict[str, Any]]],
    trace_findings: List[Dict[str, Any]],
    configs: List[SequenceStepConfig],
    lay_metadata: Dict[str, Dict[str, Any]],
) -> List[Dict[str, Any]]:
    summary = []
    layout_summary = summarize_layout_format(layout_matches)
    expected_format = layout_summary.get("inferred_plate_format", "unknown")

    for config in configs:
        seq_meta = get_sequence_labware_metadata(
            config.sequence,
            lay_metadata,
            getattr(config, "sequence_count", 0),
        )

        sequence_format = seq_meta.get("labware_format", "unknown")
        labware_name = seq_meta.get("labware_name", "unknown")

        status = "pass"
        comparison = "Not compared to uploaded destination layout."

        if config.role == "dispense":
            comparison = f"Uploaded layout appears to be {expected_format}."

            if expected_format != "unknown" and sequence_format != "unknown" and sequence_format != expected_format:
                status = "fail"

        if config.role == "aspirate":
            comparison = "Aspirate/source labware is reported, but not compared to destination layout format."

        if config.role == "tip pick up":
            comparison = "Tip-pick labware is reported for tip type and volume rule checking."

        message = (
            f"{config.role.title()} sequence '{config.sequence}' resolved to labware "
            f"'{labware_name}' from method/deck metadata. Detected format/type: {sequence_format}. {comparison}"
        )

        if status == "fail":
            message = (
                f"URGENT ERROR: Destination labware mismatch. Dispense sequence '{config.sequence}' "
                f"uses labware '{labware_name}' detected as {sequence_format}, but uploaded "
                f"layout appears to be {expected_format}."
            )

        summary.append({
            "type": "Labware Check",
            "status": status,
            "set_order": config.set_order,
            "step_role": config.role,
            "sequence": config.sequence,
            "resolved_labware_name": labware_name,
            "detected_labware_format_or_type": sequence_format,
            "labware_file": seq_meta.get("labware_file", ""),
            "labware_template_or_carrier": seq_meta.get("labware_template", ""),
            "deck_site_id": seq_meta.get("labware_site_id", ""),
            "primary_sequence_objid": seq_meta.get("primary_objid", ""),
            "layer_labware_name": seq_meta.get("layer_labware_name", ""),
            "all_detected_labware_names": seq_meta.get("all_detected_labware_names", []),
            "selected_position_count": seq_meta.get("selected_position_count", 0),
            "position_count_note": "Selected sequence positions are context only and are not used as labware capacity proof.",
            "expected_uploaded_layout_format": expected_format if config.role == "dispense" else "not applicable",
            "method_metadata_basis": seq_meta.get("labware_format_basis", ""),
            "message": message,
        })

    return summary


def validate_layout_matches(
    layout_matches: Dict[str, List[Dict[str, Any]]],
    pairs: List[TransferPair],
    marker: str,
) -> Tuple[List[Dict[str, Any]], List[ValidationIssue]]:
    findings = []
    issues = []

    marker = str(marker or "").upper().strip()

    all_layout_positions = []
    position_to_markers = defaultdict(list)
    marker_replicate_counts = {}

    for marker_label, entries in layout_matches.items():
        normalized_positions = []

        for entry in entries:
            pos = normalize_position_text(entry.get("dest_value", ""))
            if not pos:
                continue

            all_layout_positions.append(pos)
            position_to_markers[pos].append(marker_label)
            normalized_positions.append(pos)

        marker_replicate_counts[marker_label] = len(normalized_positions)

    true_duplicate_wells = {
        pos: sorted(set(labels))
        for pos, labels in position_to_markers.items()
        if len(set(labels)) > 1
    }

    for pos, labels in true_duplicate_wells.items():
        issues.append(ValidationIssue(
            "major",
            "layout",
            f"Layout well {pos} contains more than one distinct marker/sample: {labels}",
            {"well": pos, "markers": labels},
        ))

    pair_positions = [
        normalize_position_text(pair.destination_position)
        for pair in pairs
        if pair.destination_position
    ]

    missing_from_pairs = sorted(set(all_layout_positions) - set(pair_positions))

    for pos in missing_from_pairs:
        issues.append(ValidationIssue(
            "major",
            "layout",
            f"Layout destination {pos} is present in uploaded layout but not represented in deterministic transfer pairs.",
        ))

    duplicate_pair_positions = sorted(pos for pos in set(pair_positions) if pair_positions.count(pos) > 1)

    for pos in duplicate_pair_positions:
        issues.append(ValidationIssue(
            "major",
            "mapping",
            f"Transfer plan uses destination well {pos} more than once.",
        ))

    expected_marker_labels = {pair.marker_label for pair in pairs if pair.marker_label}
    actual_marker_labels = set(layout_matches.keys())

    findings.append({
        "type": "Layout Summary",
        "marker_prefix": marker,
        "sample_marker_count": len(layout_matches),
        "total_layout_destinations": len(all_layout_positions),
        "unique_destination_wells": len(set(all_layout_positions)),
        "replicate_counts_by_marker": marker_replicate_counts,
        "markers_with_replicates": {
            label: count
            for label, count in marker_replicate_counts.items()
            if count > 1
        },
        "replicate_explanation": (
            "Repeated occurrences of the same marker are treated as replicates. "
            "They are not duplicate errors unless different markers share the same exact well."
        ),
        "true_duplicate_wells": true_duplicate_wells,
        "missing_layout_wells_from_transfer_plan": missing_from_pairs,
        "extra_layout_markers_not_in_transfer_plan": sorted(actual_marker_labels - expected_marker_labels),
    })

    return findings, issues



def get_latest_trc_file(
    run_start_time: Optional[float] = None,
    wait_seconds: int = 45,
    trace_dir: str = DEFAULT_TRACE_DIR,
    require_after_start: bool = False,
) -> Optional[str]:
    pattern = os.path.join(trace_dir, "STAR_OEM_noFan_*_Trace.trc")
    end_time = time.time() + wait_seconds
    latest_any = None

    while time.time() < end_time:
        trc_files = glob.glob(pattern)

        if trc_files:
            trc_files.sort(key=os.path.getmtime, reverse=True)
            latest_any = trc_files[0]

            if run_start_time is None:
                return latest_any

            after_start = [
                f for f in trc_files
                if os.path.getmtime(f) >= (run_start_time - 10)
            ]

            if after_start:
                after_start.sort(key=os.path.getmtime, reverse=True)
                return after_start[0]

            if not require_after_start:
                return latest_any

        time.sleep(1)

    return latest_any if not require_after_start else None


def parse_trc_file(trc_path: str) -> Tuple[List[str], List[str], List[Dict[str, Any]]]:
    aspirate_steps = []
    dispense_steps = []
    ordered_steps = []

    patterns = [
        ("tip_pick", re.compile(r"1000\s*ul\s+channel\s+tip\s+pick\s+up\s+\(single step\)\s*-\s*complete[:;]\s*(.*)",
                                re.IGNORECASE)),
        ("aspirate",
         re.compile(r"1000\s*ul\s+channel\s+aspirate\s+\(single step\)\s*-\s*complete[:;]\s*(.*)", re.IGNORECASE)),
        ("dispense",
         re.compile(r"1000\s*ul\s+channel\s+dispense\s+\(single step\)\s*-\s*complete[:;]\s*(.*)", re.IGNORECASE)),
        ("tip_eject",
         re.compile(r"1000\s*ul\s+channel\s+tip\s+eject\s+\(single step\)\s*-\s*complete[:;]\s*(.*)", re.IGNORECASE)),
        ("aspirate_96_head",
         re.compile(r"co-?re\s+96\s+head\s+aspirate\s+\(single step\)\s*-\s*complete[:;]\s*(.*)", re.IGNORECASE)),
        ("dispense_96_head",
         re.compile(r"co-?re\s+96\s+head\s+dispense\s+\(single step\)\s*-\s*complete[:;]\s*(.*)", re.IGNORECASE)),
    ]

    with open(trc_path, "r", encoding="latin-1") as file:
        lines = file.readlines()

    order = 1

    for raw_line in lines:
        line = raw_line.strip()

        for step_type, pattern in patterns:
            match = pattern.search(line)
            if not match:
                continue

            payload = match.group(1).strip()

            if step_type in ("aspirate", "aspirate_96_head"):
                aspirate_steps.append(payload)

            if step_type in ("dispense", "dispense_96_head"):
                dispense_steps.append(payload)

            ordered_steps.append({
                "order": order,
                "type": step_type,
                "raw": payload,
            })

            order += 1
            break

    return aspirate_steps, dispense_steps, ordered_steps


def validate_expected_runtime_steps(
    execution_plan: List[Dict[str, Any]],
    ordered_steps: List[Dict[str, Any]],
    trace_text: str,
) -> Tuple[List[Dict[str, Any]], List[ValidationIssue]]:
    findings = []
    issues = []

    expected_actions = [
        item.get("action")
        for item in execution_plan
        if item.get("action") in {"tip_pick", "aspirate", "dispense", "tip_eject"}
    ]

    actual_actions = []

    for step in ordered_steps:
        step_type = step.get("type", "")

        if "aspirate" in step_type:
            actual_actions.append("aspirate")
        elif "dispense" in step_type:
            actual_actions.append("dispense")
        elif "tip_pick" in step_type or "tip pick" in step_type:
            actual_actions.append("tip_pick")
        elif "tip_eject" in step_type or "tip eject" in step_type:
            actual_actions.append("tip_eject")

    trace_lower = trace_text.lower()

    runtime_error_reason = ""

    error_match = re.search(
        r"main\s*-\s*error;\s*(.*?)(?:\n|$)",
        trace_text,
        re.IGNORECASE,
    )

    if error_match:
        runtime_error_reason = error_match.group(1).strip()

    if "tip type mismatch" in trace_lower:
        runtime_error_reason = "Tip type mismatch. The selected liquid class is not compatible with the picked tip type."

    expected_count = len(expected_actions)
    actual_count = len(actual_actions)

    status = "pass" if actual_count >= expected_count and not runtime_error_reason else "fail"

    finding = {
        "type": "expected_runtime_step_check",
        "status": status,
        "expected_step_count": expected_count,
        "actual_step_count": actual_count,
        "expected_actions": expected_actions,
        "actual_actions": actual_actions,
        "runtime_error_reason": runtime_error_reason,
    }

    findings.append(finding)

    if status != "pass":
        message = (
            f"Runtime did not execute all planned steps. "
            f"Expected {expected_count} liquid-handling steps, but trace showed {actual_count}."
        )

        if runtime_error_reason:
            message += f" Reason: {runtime_error_reason}"

        issues.append(ValidationIssue(
            "critical",
            "runtime_step_execution",
            message,
            finding,
        ))

    return findings, issues


def normalize_channel_steps(step_lines: List[str]) -> List[Tuple[str, str, str, str]]:
    normalized = []

    pattern = re.compile(
        r">\s*channel\s+(\d+):\s*([^,]+),\s*([A-Za-z]*\d+)\s+([\d\.]+)\s*[µu]L",
        re.IGNORECASE,
    )

    fallback_pattern = re.compile(
        r"channel\s+(\d+).*?([A-Za-z0-9_\- ]+).*?\b([A-H]\d{1,2}|\d+)\b.*?([\d\.]+)\s*[µu]L",
        re.IGNORECASE,
    )

    for step_line in step_lines:
        matches = pattern.findall(step_line)

        if not matches:
            matches = fallback_pattern.findall(step_line)

        for channel_num, labware, position, volume in matches:
            normalized.append((
                channel_num.strip(),
                labware.strip(),
                normalize_position_text(position),
                volume.strip(),
            ))

    return normalized


def cross_match_positions(aspirate_steps: List[str], dispense_steps: List[str], replicate_value: str) -> Tuple[bool, Any]:
    replicate_count = REPLICATE_MAP.get(str(replicate_value).lower(), 1)

    normalized_aspirate_steps = normalize_channel_steps(aspirate_steps)
    normalized_dispense_steps = normalize_channel_steps(dispense_steps)

    aspirate_channel_map = {}
    for aspirate in normalized_aspirate_steps:
        aspirate_channel_map[aspirate[0]] = aspirate

    aspirate_to_dispense_map = defaultdict(list)

    for dispense in normalized_dispense_steps:
        channel = dispense[0]

        if channel in aspirate_channel_map:
            aspirate_step = aspirate_channel_map[channel]
            aspirate_to_dispense_map[aspirate_step].append(dispense[1:])
        else:
            return False, f"Channel {channel} in dispense step does not match any aspirate step."

    for aspirate_step, dispenses in aspirate_to_dispense_map.items():
        if replicate_count > 0 and len(dispenses) % replicate_count != 0:
            return False, (
                f"Aspirate step {aspirate_step} has {len(dispenses)} dispenses, "
                f"which is not divisible by replicate count {replicate_count}."
            )

    return True, dict(aspirate_to_dispense_map)


def validate_trace_against_layout(aspirate_to_dispense_map: Dict[Any, Any], layout_matches: Dict[str, List[Dict[str, Any]]], marker: str) -> Tuple[List[Dict[str, Any]], List[ValidationIssue]]:
    findings = []
    issues = []
    marker = marker.upper().strip()

    expected_by_sample_num = defaultdict(list)

    for marker_key, entries in layout_matches.items():
        try:
            sample_num = int(marker_key.upper().replace(marker, ""))
            expected_by_sample_num[sample_num].extend(entries)
        except ValueError:
            issues.append(ValidationIssue("minor", "marker", f"Could not parse sample number from marker {marker_key}"))

    seen_trace_destinations = []

    for aspirate_step, dispenses in aspirate_to_dispense_map.items():
        source_position_text = str(aspirate_step[2]).strip()
        source_num = extract_position_number(source_position_text)

        if source_num is None:
            issues.append(ValidationIssue("major", "trace", f"Could not extract source number from aspirate position: {source_position_text}"))
            continue

        marker_label = f"{marker}{source_num}"
        expected_entries = expected_by_sample_num.get(source_num, [])
        actual_positions = [normalize_position_text(dispense[1]) for dispense in dispenses]
        seen_trace_destinations.extend(actual_positions)

        expected_positions = [
            normalize_position_text(entry.get("dest_value", ""))
            for entry in expected_entries
            if entry.get("dest_value") not in [None, "", "nan"]
        ]

        missing = sorted(set(expected_positions) - set(actual_positions))
        unexpected = sorted(set(actual_positions) - set(expected_positions))

        findings.append({
            "type": "trace_layout_match",
            "marker": marker_label,
            "source_position": source_position_text,
            "expected_destinations": expected_positions,
            "actual_destinations": actual_positions,
            "missing_destinations": missing,
            "unexpected_destinations": unexpected,
            "status": "pass" if not missing and not unexpected else "fail",
        })

        if missing:
            issues.append(ValidationIssue("critical", "trace_layout", f"{marker_label} missing expected destination wells: {missing}"))

        if unexpected:
            issues.append(ValidationIssue("critical", "trace_layout", f"{marker_label} has unexpected destination wells: {unexpected}"))

        if len(actual_positions) != len(expected_positions):
            issues.append(ValidationIssue("major", "trace_layout", f"{marker_label} source/destination count mismatch", {
                "expected_count": len(expected_positions),
                "actual_count": len(actual_positions),
            }))

    duplicate_trace_dests = sorted([pos for pos in set(seen_trace_destinations) if seen_trace_destinations.count(pos) > 1])
    for pos in duplicate_trace_dests:
        issues.append(ValidationIssue("major", "trace", f"Duplicate destination well found in trace: {pos}"))

    all_expected = []
    for entries in layout_matches.values():
        all_expected.extend([normalize_position_text(entry.get("dest_value", "")) for entry in entries])

    missing_global = sorted(set(all_expected) - set(seen_trace_destinations))
    for pos in missing_global:
        issues.append(ValidationIssue("major", "trace_layout", f"Layout destination not represented in trace: {pos}"))

    labware_findings, labware_issues = validate_trace_labware_against_layout_format(
        aspirate_to_dispense_map,
        layout_matches,
    )
    findings.extend(labware_findings)
    issues.extend(labware_issues)

    return findings, issues


def validate_tracking_source_order(aspirate_to_dispense_map: Dict[Any, Any], marker: str) -> Tuple[List[Dict[str, Any]], List[ValidationIssue]]:
    findings = []
    issues = []
    marker = marker.upper().strip()

    expected_source = 1

    for aspirate_step, dispenses in aspirate_to_dispense_map.items():
        source_position_text = str(aspirate_step[2]).strip()
        source_num = extract_position_number(source_position_text)

        if source_num is None:
            issues.append(ValidationIssue("major", "tracking", f"Could not extract source number from aspirate position: {source_position_text}"))
            continue

        marker_label = f"{marker}{source_num}"
        status = "pass" if source_num == expected_source else "fail"

        findings.append({
            "type": "tracking_source_order",
            "marker": marker_label,
            "expected_source_number": expected_source,
            "actual_source_number": source_num,
            "source_position": source_position_text,
            "dispense_count": len(dispenses),
            "status": status,
        })

        if status != "pass":
            issues.append(ValidationIssue(
                "major",
                "tracking",
                f"Tracking order mismatch. Expected sample/source {expected_source}, found {source_num}.",
                {"marker": marker_label, "source_position": source_position_text},
            ))

        expected_source += 1

    return findings, issues


def count_active_channels(channel_pattern: str) -> int:
    text = str(channel_pattern or "").strip()
    return sum(1 for char in text if char == "1")


def infer_transfer_direction_from_positions(positions: List[str]) -> str:
    normalized = [normalize_position_text(pos) for pos in positions if pos]

    parsed = []
    for pos in normalized:
        match = re.match(r"^([A-P])(\d{1,2})$", pos)
        if match:
            parsed.append((match.group(1), int(match.group(2)), pos))
        else:
            num = extract_position_number(pos)
            if num is not None:
                parsed.append(("", num, pos))

    if len(parsed) < 2:
        return "single position / direction not inferable"

    rows = [item[0] for item in parsed if item[0]]
    cols = [item[1] for item in parsed]

    if rows and len(set(rows)) == 1 and len(set(cols)) > 1:
        return "left-to-right across columns"

    if rows and len(set(cols)) == 1 and len(set(rows)) > 1:
        return "top-to-bottom down rows"

    if cols == sorted(cols):
        return "ascending numeric order"

    if cols == sorted(cols, reverse=True):
        return "descending numeric order"

    return "mixed/non-linear order"


def infer_labware_orientation_from_layout(layout_matches: Dict[str, List[Dict[str, Any]]]) -> Dict[str, Any]:
    layout_summary = summarize_layout_format(layout_matches)
    row_count = int(layout_summary.get("row_count") or 0)
    column_count = int(layout_summary.get("column_count") or 0)

    if column_count > row_count:
        orientation = "landscape"
    elif row_count > column_count:
        orientation = "portrait"
    else:
        orientation = "square/unknown"

    return {
        "type": "labware_orientation",
        "orientation": orientation,
        "row_count": row_count,
        "column_count": column_count,
        "inferred_plate_format": layout_summary.get("inferred_plate_format", "unknown"),
    }


def analyze_orientation_and_tip_usage(
    configs: List[SequenceStepConfig],
    execution_plan: List[Dict[str, Any]],
    layout_matches: Dict[str, List[Dict[str, Any]]],
    aspirate_steps: List[str],
    dispense_steps: List[str],
    cross_match: Any,
) -> Tuple[List[Dict[str, Any]], List[ValidationIssue]]:
    findings = []
    issues = []

    orientation = infer_labware_orientation_from_layout(layout_matches) if layout_matches else {
        "type": "labware_orientation",
        "orientation": "unknown/no layout",
        "row_count": 0,
        "column_count": 0,
        "inferred_plate_format": "unknown",
    }

    findings.append(orientation)

    normalized_aspirates = normalize_channel_steps(aspirate_steps)
    normalized_dispenses = normalize_channel_steps(dispense_steps)

    aspirate_channels = sorted({item[0] for item in normalized_aspirates}, key=lambda x: int(x) if str(x).isdigit() else 999)
    dispense_channels = sorted({item[0] for item in normalized_dispenses}, key=lambda x: int(x) if str(x).isdigit() else 999)

    aspirate_positions = [item[2] for item in normalized_aspirates]
    dispense_positions = [item[2] for item in normalized_dispenses]

    findings.append({
        "type": "channel_transfer_direction",
        "labware_orientation": orientation.get("orientation"),
        "aspirate_channel_count": len(aspirate_channels),
        "dispense_channel_count": len(dispense_channels),
        "aspirate_channels_used": aspirate_channels,
        "dispense_channels_used": dispense_channels,
        "aspirate_direction": infer_transfer_direction_from_positions(aspirate_positions),
        "dispense_direction": infer_transfer_direction_from_positions(dispense_positions),
        "transfer_mode": "single tip/channel at a time" if max(len(aspirate_channels), len(dispense_channels)) <= 1 else "multiple tips/channels in parallel",
    })

    tip_pick_steps = [item for item in execution_plan if item.get("action") == "tip_pick"]

    for index, tip_step in enumerate(tip_pick_steps, start=1):
        pattern = str(tip_step.get("channel_pattern") or "")
        picked_tip_count = count_active_channels(pattern)

        if str(tip_step.get("hardware_mode", "")).lower() == "head":
            picked_tip_count = 96

        used_aspirate_count = len(aspirate_channels)
        used_dispense_count = len(dispense_channels)
        max_used_count = max(used_aspirate_count, used_dispense_count)

        unused_count = max(picked_tip_count - max_used_count, 0)

        status = "pass"
        if picked_tip_count > 0 and max_used_count == 0:
            status = "fail"
        elif unused_count > 0:
            status = "warning"

        finding = {
            "type": "tip_usage_check",
            "tip_pick_step_index": index,
            "tip_sequence": tip_step.get("sequence"),
            "hardware_mode": tip_step.get("hardware_mode"),
            "channel_pattern": pattern,
            "picked_tip_count": picked_tip_count,
            "aspirate_channels_used_count": used_aspirate_count,
            "dispense_channels_used_count": used_dispense_count,
            "unused_picked_tip_count": unused_count,
            "aspirate_channels_used": aspirate_channels,
            "dispense_channels_used": dispense_channels,
            "status": status,
            "message": (
                f"Picked up {picked_tip_count} tip(s); used "
                f"{used_aspirate_count} aspirate channel(s) and "
                f"{used_dispense_count} dispense channel(s)."
            ),
        }

        findings.append(finding)

        if status == "fail":
            issues.append(ValidationIssue(
                "critical",
                "tip_usage",
                f"Tips were picked up from {tip_step.get('sequence')}, but no aspirate/dispense channel usage was detected.",
                finding,
            ))

        elif status == "warning":
            issues.append(ValidationIssue(
                "major",
                "tip_usage",
                (
                    f"Possible unused tips: {tip_step.get('sequence')} picked up "
                    f"{picked_tip_count} tip(s), but only {max_used_count} channel(s) were used."
                ),
                finding,
            ))

    return findings, issues


def generate_pyhamilton_review_script(
    configs: List[SequenceStepConfig],
    execution_plan: List[Dict[str, Any]],
    output_dir: str,
) -> str:
    os.makedirs(output_dir, exist_ok=True)

    script_path = os.path.join(
        output_dir,
        f"devpal_lite_review_required_pyhamilton_script_{datetime.now().strftime('%H%M%S')}.py",
    )

    payload = {
        "review_required": True,
        "generated_at": datetime.now().isoformat(),
        "selected_sequences": [asdict(config) for config in configs],
        "execution_plan": execution_plan,
    }

    lines = [
        "from pyhamilton import *",
        "import inspect",
        "",
        "REVIEW_REQUIRED = True",
        f"DEV_PAL_LITE_PLAN = {json.dumps(payload, indent=4)}",
        "",
        "def _normalize_hardware_mode(value):",
        "    text = str(value or '').strip().lower().replace(' ', '').replace('_', '-').replace('-', '')",
        "    if text in {'head384', '384head', '384', 'mph384', '384mph'}:",
        "        return 'head384'",
        "    if text in {'head96', '96head', 'head', '96', 'mph96', '96mph'}:",
        "        return 'head96'",
        "    return 'channels'",
        "",
        "def _head_size(mode):",
        "    mode = _normalize_hardware_mode(mode)",
        "    if mode == 'head384':",
        "        return 384",
        "    if mode == 'head96':",
        "        return 96",
        "    return 0",
        "",
        "def _get_function(*names):",
        "    for name in names:",
        "        func = globals().get(name)",
        "        if func is not None:",
        "            return func",
        "    return None",
        "",
        "def _call_best_effort(func, *args, **kwargs):",
        "    if func is None:",
        "        raise RuntimeError('Required PyHamilton function is unavailable.')",
        "    try:",
        "        signature = inspect.signature(func)",
        "        allowed = {k: v for k, v in kwargs.items() if k in signature.parameters}",
        "        return func(*args, **allowed)",
        "    except TypeError:",
        "        return func(*args)",
        "",
        "def _resolve_resource_class(labware_name):",
        "    text = str(labware_name or '').upper()",
        "    preferred_names = []",
        "    if '384' in text:",
        "        preferred_names = ['Plate384', 'Plate96']",
        "    elif '24' in text:",
        "        preferred_names = ['Plate24', 'Plate96']",
        "    elif '96' in text or 'PLATE' in text:",
        "        preferred_names = ['Plate96']",
        "    elif 'TIP' in text or 'TIPRACK' in text or 'SLIM' in text or 'HTF' in text:",
        "        preferred_names = ['Tip96', 'TipRack', 'Plate96']",
        "    elif 'RACK' in text or 'TUBE' in text or 'HONEYCOMB' in text or 'CARRIER' in text:",
        "        preferred_names = ['Plate96', 'Tip96']",
        "    else:",
        "        preferred_names = ['Plate96']",
        "    for class_name in preferred_names:",
        "        if class_name in globals():",
        "            return globals()[class_name]",
        "    return Plate96",
        "",
        "def _assign_sequence_resource(layout_manager, labware_name):",
        "    labware_name = str(labware_name or '').strip()",
        "    if not labware_name:",
        "        return None",
        "    resource_class = _resolve_resource_class(labware_name)",
        "    attempts = []",
        "    if hasattr(layout_manager, 'assign_resource'):",
        "        attempts.append(lambda: layout_manager.assign_resource(labware_name))",
        "    if 'ResourceType' in globals() and resource_class is not None:",
        "        if hasattr(layout_manager, 'assign_resource'):",
        "            attempts.append(lambda: layout_manager.assign_resource(ResourceType(resource_class, labware_name)))",
        "        if hasattr(layout_manager, 'assign_unused_resource'):",
        "            attempts.append(lambda: layout_manager.assign_unused_resource(ResourceType(resource_class, labware_name)))",
        "            attempts.append(lambda: layout_manager.assign_unused_resource(ResourceType(resource_class, '')))",
        "    last_error = None",
        "    for attempt in attempts:",
        "        try:",
        "            resource = attempt()",
        "            if resource is not None:",
        "                return resource",
        "        except Exception as exc:",
        "            last_error = exc",
        "    raise RuntimeError(f'Could not assign labware resource for {labware_name!r}. Last error: {last_error}')",
        "",
        "def _tip_pick_head(ham_int, seq, head_size):",
        "    func = _get_function('tip_pick_up_head_seq', 'tip_pick_up_384_seq' if head_size == 384 else 'tip_pick_up_96_seq')",
        "    if func:",
        "        return _call_best_effort(func, ham_int, tip_seq=seq, tip384_seq=seq, tip96_seq=seq, sequence=seq, head_size=head_size)",
        "    command = globals().get('PICKUP384' if head_size == 384 else 'PICKUP96', 'PICKUP384' if head_size == 384 else 'PICKUP96')",
        "    tid = ham_int.send_command(command, tipSequence=seq, channelVariable='1' * head_size, sequenceCounting=1)",
        "    return ham_int.wait_on_response(tid, raise_first_exception=True, timeout=120)",
        "",
        "def _aspirate_head(ham_int, resource, seq, volume, liquid_class, head_size):",
        "    func = _get_function('aspirate_head_seq', 'aspirate_384_seq' if head_size == 384 else 'aspirate_96_seq')",
        "    if func:",
        "        return _call_best_effort(func, ham_int, resource, plate=resource, plate384=resource, plate96=resource, head_asp_seq=seq, asp_seq=seq, sequence=seq, vols=volume, volume=volume, liq_class=liquid_class, liquid_class=liquid_class, liq_class2=liquid_class, head_size=head_size)",
        "    command = globals().get('ASPIRATE384' if head_size == 384 else 'ASPIRATE96', 'ASPIRATE384' if head_size == 384 else 'ASPIRATE96')",
        "    tid = ham_int.send_command(command, aspirateSequence=seq, labwarePositions='', aspirateVolume=volume, volumes=volume, channelVariable='1' * head_size, liquidClass=liquid_class, sequenceCounting=0, capacitiveLLD=2)",
        "    return ham_int.wait_on_response(tid, raise_first_exception=True, timeout=120)",
        "",
        "def _dispense_head(ham_int, resource, seq, volume, liquid_class, head_size):",
        "    func = _get_function('dispense_head_seq', 'dispense_384_seq' if head_size == 384 else 'dispense_96_seq2', 'dispense_96_seq')",
        "    if func:",
        "        return _call_best_effort(func, ham_int, resource, plate=resource, plate384=resource, plate96=resource, head_disp_seq=seq, disp_seq=seq, sequence=seq, vols=volume, volume=volume, liq_class=liquid_class, liquid_class=liquid_class, liq_class2=liquid_class, head_size=head_size)",
        "    command = globals().get('DISPENSE384' if head_size == 384 else 'DISPENSE96', 'DISPENSE384' if head_size == 384 else 'DISPENSE96')",
        "    tid = ham_int.send_command(command, dispenseSequence=seq, labwarePositions='', dispenseVolume=volume, volumes=volume, channelVariable='1' * head_size, liquidClass=liquid_class, sequenceCounting=0)",
        "    return ham_int.wait_on_response(tid, raise_first_exception=True, timeout=120)",
        "",
        "def _eject_head(ham_int, head_size):",
        "    func = _get_function('tip_eject_head', 'tip_eject_384' if head_size == 384 else 'tip_eject_96')",
        "    if func:",
        "        return _call_best_effort(func, ham_int, head_size=head_size)",
        "    command = globals().get('EJECT384' if head_size == 384 else 'EJECT96', 'EJECT384' if head_size == 384 else 'EJECT96')",
        "    tid = ham_int.send_command(command, labwarePositions='', channelVariable='1' * head_size, tipEjectToKnownPosition=2)",
        "    return ham_int.wait_on_response(tid, raise_first_exception=True, timeout=120)",
        "",
        "def run_review_required_plan(ham_int, layout_manager):",
        "    if REVIEW_REQUIRED:",
        "        raise RuntimeError('Review required before running generated PyHamilton script. Set REVIEW_REQUIRED = False only after human review.')",
        "    resource_cache = {}",
        "    initialize(ham_int)",
    ]

    for item in execution_plan:
        action = str(item.get("action") or "").strip()
        seq = str(item.get("sequence") or "").strip()
        hardware_mode = normalize_hardware_mode(item.get("hardware_mode") or "channels")
        head_size = get_head_size_from_mode(hardware_mode)
        channel = normalize_assigned_channel_pattern(item.get("channel_pattern") or "", hardware_mode)
        volume = float(item.get("volume_ul") or 0)
        liquid_class = str(item.get("liquid_class") or "").strip()
        increment = int(item.get("increment") or 0)
        labware_name = str(item.get("labware_name") or seq).strip()
        cache_key = str(labware_name or seq).strip().lower()

        lines.append("")
        lines.append(f"    print({('RUNNING order=' + str(item.get('order')) + ' set=' + str(item.get('set_order')) + ' action=' + action + ' sequence=' + seq + ' hardware=' + hardware_mode)!r})")

        if (
            item.get("autoincrement")
            and action in {"aspirate", "dispense"}
            and not bool(item.get("manual"))
            and seq
            and increment > 0
        ):
            lines.append(f"    inc_sequence(ham_int, sequence={seq!r}, increment={increment!r})")

        if action == "tip_pick":
            if head_size:
                lines.append(f"    _tip_pick_head(ham_int, {seq!r}, {head_size!r})")
            else:
                lines.append(f"    tip_pick_up_seq(ham_int, tipseq={seq!r}, channel={channel!r})")

        elif action == "aspirate":
            if head_size:
                lines.append(f"    if {cache_key!r} not in resource_cache:")
                lines.append(f"        resource_cache[{cache_key!r}] = _assign_sequence_resource(layout_manager, {labware_name!r})")
                lines.append(f"    _aspirate_head(ham_int, resource_cache[{cache_key!r}], {seq!r}, {volume!r}, {liquid_class!r}, {head_size!r})")
            else:
                lines.append(f"    aspirate_seq(ham_int, asp_seq={seq!r}, vols={volume!r}, channel={channel!r}, liq_class={liquid_class!r})")

        elif action == "dispense":
            if head_size:
                lines.append(f"    if {cache_key!r} not in resource_cache:")
                lines.append(f"        resource_cache[{cache_key!r}] = _assign_sequence_resource(layout_manager, {labware_name!r})")
                lines.append(f"    _dispense_head(ham_int, resource_cache[{cache_key!r}], {seq!r}, {volume!r}, {liquid_class!r}, {head_size!r})")
            else:
                lines.append(f"    dispense_seq(ham_int, disp_seq={seq!r}, vols={volume!r}, channel={channel!r}, liq_class={liquid_class!r})")

        elif action == "tip_eject":
            if head_size:
                lines.append(f"    _eject_head(ham_int, {head_size!r})")
            else:
                lines.append(f"    tip_eject_seq2(ham_int, waste_seq={seq!r}, channel={channel!r})")

    Path(script_path).write_text("\n".join(lines) + "\n", encoding="utf-8")
    return script_path

def humanize_key(key: str) -> str:
    return str(key).replace("_", " ").replace("-", " ").title()


def humanize_type(value: Any) -> str:
    return humanize_key(str(value or "")).strip()


def sort_position_key(value: Any):
    text = str(value or "").strip().upper()
    match = re.match(r"^([A-Z]*)(\d+)$", text)
    if match:
        return (match.group(1), int(match.group(2)))
    return (text, 0)


def is_error_record(value: Any) -> bool:
    text = str(value).lower()
    return (
        "urgent" in text
        or "critical" in text
        or "error" in text
        or "fail" in text
        or "mismatch" in text
        or "wrong labware" in text
    )


def safe_text(value: Any) -> str:
    if value is None:
        return "None"
    return str(value)


def format_report_value(value: Any) -> str:
    if isinstance(value, dict):
        if not value:
            return "None"
        return "; ".join(
            f"{humanize_key(key)}: {format_report_value(val)}"
            for key, val in value.items()
        )

    if isinstance(value, list):
        if not value:
            return "None"
        if len(value) > 30:
            shown = ", ".join(safe_text(item) for item in value[:30])
            return f"{shown} ... ({len(value)} total)"
        return ", ".join(safe_text(item) for item in value)

    return safe_text(value)


def readable_record_title(section_title: str, item: Any, index: int) -> str:
    section = str(section_title or "").lower()

    if not isinstance(item, dict):
        return f"{section_title} {index}"

    if "selected sequence" in section:
        return (
            f"Sequence: {item.get('sequence', 'Unknown')} | "
            f"Role: {item.get('role', item.get('step_role', 'Unknown'))} | "
            f"Set: {item.get('set_order', 'N/A')}"
        )

    if "execution plan" in section:
        return (
            f"Plan Step {item.get('order', index)}: "
            f"{humanize_type(item.get('action', 'Step'))} | "
            f"{item.get('sequence', 'Unknown')} | "
            f"Set {item.get('set_order', 'N/A')}"
        )

    if "transfer pair" in section:
        return (
            f"{item.get('pair_id', f'Transfer Pair {index}')}: "
            f"{item.get('source_sequence', '')} → {item.get('destination_sequence', '')}"
        )

    if "layout finding" in section:
        return humanize_type(item.get("type", f"Layout Finding {index}"))

    if "trace finding" in section:
        return humanize_type(item.get("type", f"Trace Finding {index}"))

    if "validation" in section or "issue" in section:
        return (
            f"{humanize_type(item.get('severity', 'Issue'))} | "
            f"{humanize_type(item.get('category', 'Validation'))}"
        )

    return f"{section_title} {index}"


def docx_add_line(doc, text: str, force_red: bool = False) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(str(text))

    if RGBColor is not None and (force_red or is_error_record(text)):
        run.font.color.rgb = RGBColor(255, 0, 0)


def write_human_section_docx(doc, title: str, data: Any) -> None:
    doc.add_heading(title, level=1)

    if not data:
        doc.add_paragraph("None found.")
        return

    if isinstance(data, dict):
        for key, value in data.items():
            docx_add_line(doc, f"{humanize_key(key)}: {format_report_value(value)}")
        return

    if isinstance(data, list):
        for index, item in enumerate(data, start=1):
            heading = readable_record_title(title, item, index)
            docx_add_line(doc, heading, force_red=is_error_record(item))

            if isinstance(item, dict):
                for key, value in item.items():
                    if key == "type":
                        continue
                    docx_add_line(doc, f"{humanize_key(key)}: {format_report_value(value)}")
            else:
                docx_add_line(doc, safe_text(item))
        return

    docx_add_line(doc, safe_text(data))


def write_human_text_section(handle, title: str, data: Any) -> None:
    handle.write(f"\n{title}\n")
    handle.write("-" * len(title) + "\n")

    if not data:
        handle.write("None found.\n")
        return

    if isinstance(data, dict):
        for key, value in data.items():
            handle.write(f"{humanize_key(key)}: {format_report_value(value)}\n")
        return

    if isinstance(data, list):
        for index, item in enumerate(data, start=1):
            handle.write(f"\n{readable_record_title(title, item, index)}\n")

            if isinstance(item, dict):
                for key, value in item.items():
                    if key == "type":
                        continue
                    handle.write(f"{humanize_key(key)}: {format_report_value(value)}\n")
            else:
                handle.write(f"{safe_text(item)}\n")
        return

    handle.write(f"{safe_text(data)}\n")

def write_docx_report(output: DevPalLiteOutput, output_dir: str) -> str:
    if Document is None:
        return write_text_report(output, output_dir)

    path = os.path.join(
        output_dir,
        f"DevPal_Lite_Validation_Report_{datetime.now().strftime('%H%M%S')}.docx",
    )

    doc = Document()

    doc.add_heading("DevPal Lite Validation Report", 0)

    summary_items = [
        ("Generated", datetime.now().isoformat()),
        ("LAY File", output.lay_file or "Not provided"),
        ("Layout File", output.layout_file or "Not provided"),
        ("Trace File", output.trace_file or "Not found"),
        ("Generated Review Script", output.generated_review_script or "Not generated"),
        ("Report Path", output.report_path or "Pending"),
    ]

    doc.add_heading("Run Summary", level=1)
    for label, value in summary_items:
        paragraph = doc.add_paragraph()
        paragraph.add_run(f"{label}: ").bold = True
        paragraph.add_run(str(value))

    urgent_issues = [
        issue for issue in output.issues
        if is_error_record(issue)
    ]

    if urgent_issues:
        doc.add_heading("Urgent / Critical Issues Found", level=1)

        for index, issue in enumerate(urgent_issues, start=1):
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(f"Issue {index}")
            run.bold = True
            if RGBColor is not None:
                run.font.color.rgb = RGBColor(255, 0, 0)

            if isinstance(issue, dict):
                for key, value in issue.items():
                    line = f"{humanize_key(key)}: {safe_text(value)}"
                    docx_add_line(doc, line, force_red=True)
            else:
                docx_add_line(doc, safe_text(issue), force_red=True)

    sections = [
        ("Selected Sequences", output.selected_sequences),
        ("Source Positions", output.source_positions),
        ("Destination Positions", output.destination_positions),
        ("Transfer Pairs", output.transfer_pairs),
        ("Execution Plan", output.execution_plan),
        ("Marker Matches", output.marker_matches),
        ("Layout Findings", output.layout_findings),
        ("Trace Findings", output.trace_findings),
        ("Validation Findings", output.validation_findings),
        ("Issues", output.issues),
        ("Limitations", output.limitations),
    ]

    for title, data in sections:
        write_human_section_docx(doc, title, data)

    doc.save(path)
    return path


def write_pdf_report(output: DevPalLiteOutput, output_dir: str) -> str:
    if SimpleDocTemplate is None:
        return write_text_report(output, output_dir)

    path = os.path.join(
        output_dir,
        f"DevPal_Lite_Validation_Report_{datetime.now().strftime('%H%M%S')}.pdf",
    )

    styles = getSampleStyleSheet()
    story = []

    def pdf_line(text: Any, style_name: str = "BodyText"):
        line = safe_text(text)

        if is_error_record(line):
            line = f'<font color="red">{line}</font>'

        story.append(Paragraph(line, styles[style_name]))

    def pdf_section(title: str, data: Any):
        story.append(Spacer(1, 10))
        story.append(Paragraph(title, styles["Heading1"]))
        story.append(Spacer(1, 4))

        if not data:
            pdf_line("None found.")
            return

        if isinstance(data, dict):
            for key, value in data.items():
                pdf_line(f"{humanize_key(key)}: {safe_text(value)}")
            return

        if isinstance(data, list):
            for index, item in enumerate(data[:200], start=1):
                item_is_error = is_error_record(item)
                item_title = f"Item {index}"

                if item_is_error:
                    item_title = f'<font color="red">{item_title}</font>'

                story.append(Paragraph(item_title, styles["Heading3"]))

                if isinstance(item, dict):
                    for key, value in item.items():
                        pdf_line(f"{humanize_key(key)}: {safe_text(value)}")
                else:
                    pdf_line(item)

            if len(data) > 200:
                pdf_line(f"Report truncated in PDF view after 200 items. Full output model contains {len(data)} items.")

            return

        pdf_line(data)

    story.append(Paragraph("DevPal Lite Validation Report", styles["Title"]))
    story.append(Spacer(1, 12))

    summary_rows = [
        ["Generated", datetime.now().isoformat()],
        ["LAY File", output.lay_file or "Not provided"],
        ["Layout File", output.layout_file or "Not provided"],
        ["Trace File", output.trace_file or "Not found"],
        ["Generated Review Script", output.generated_review_script or "Not generated"],
    ]

    summary_table = Table(summary_rows, colWidths=[130, 380])
    summary_table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("BACKGROUND", (0, 0), (0, -1), colors.lightgrey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))

    story.append(summary_table)
    story.append(Spacer(1, 12))

    urgent_issues = [
        issue for issue in output.issues
        if is_error_record(issue)
    ]

    if urgent_issues:
        story.append(Paragraph('<font color="red">Urgent / Critical Issues Found</font>', styles["Heading1"]))
        for index, issue in enumerate(urgent_issues, start=1):
            story.append(Paragraph(f'<font color="red">Issue {index}</font>', styles["Heading3"]))

            if isinstance(issue, dict):
                for key, value in issue.items():
                    line = f"{humanize_key(key)}: {safe_text(value)}"
                    story.append(Paragraph(f'<font color="red">{line}</font>', styles["BodyText"]))
            else:
                story.append(Paragraph(f'<font color="red">{safe_text(issue)}</font>', styles["BodyText"]))

        story.append(Spacer(1, 10))

    sections = [
        ("Selected Sequences", output.selected_sequences),
        ("Source Positions", output.source_positions),
        ("Destination Positions", output.destination_positions),
        ("Transfer Pairs", output.transfer_pairs),
        ("Execution Plan", output.execution_plan),
        ("Marker Matches", output.marker_matches),
        ("Layout Findings", output.layout_findings),
        ("Trace Findings", output.trace_findings),
        ("Validation Findings", output.validation_findings),
        ("Issues", output.issues),
        ("Limitations", output.limitations),
    ]

    for title, data in sections:
        pdf_section(title, data)

    SimpleDocTemplate(path, pagesize=letter).build(story)
    return path


def write_text_report(output: DevPalLiteOutput, output_dir: str) -> str:
    path = os.path.join(output_dir, f"DevPal_Lite_Validation_Report_{datetime.now().strftime('%H%M%S')}.txt")

    with open(path, "w", encoding="utf-8") as handle:
        handle.write("DevPal Lite Validation Report\n")
        handle.write("=" * 80 + "\n")
        handle.write(f"Generated: {datetime.now().isoformat()}\n")
        handle.write(f"LAY File: {output.lay_file or 'Not provided'}\n")
        handle.write(f"Layout File: {output.layout_file or 'Not provided'}\n")
        handle.write(f"Trace File: {output.trace_file or 'Not found'}\n")
        handle.write(f"Generated Review Script: {output.generated_review_script or 'Not generated'}\n")

        sections = [
            ("Selected Sequences", output.selected_sequences),
            ("Source Positions", output.source_positions),
            ("Destination Positions", output.destination_positions),
            ("Transfer Pairs", output.transfer_pairs),
            ("Execution Plan", output.execution_plan),
            ("Marker Matches", output.marker_matches),
            ("Layout Findings", output.layout_findings),
            ("Trace Findings", output.trace_findings),
            ("Validation Findings", output.validation_findings),
            ("Issues", output.issues),
            ("Limitations", output.limitations),
        ]

        for title, data in sections:
            write_human_text_section(handle, title, data)

    return path


def create_output_model(
    lay_file: str,
    layout_file: str,
    trace_file: str,
    configs: List[SequenceStepConfig],
    pairs: List[TransferPair],
    execution_plan: List[Dict[str, Any]],
    marker_matches: Dict[str, List[Dict[str, Any]]],
    layout_findings: List[Dict[str, Any]],
    trace_findings: List[Dict[str, Any]],
    issues: List[ValidationIssue],
    limitations: List[str],
    review_script_path: str,
    report_path: str = "",
) -> DevPalLiteOutput:
    source_positions = sorted(
        {pair.source_position for pair in pairs if pair.source_position},
        key=sort_position_key,
    )

    destination_positions = sorted(
        {pair.destination_position for pair in pairs if pair.destination_position},
        key=sort_position_key,
    )

    validation_findings = [
        {
            "severity": issue.severity,
            "category": humanize_type(issue.category),
            "message": issue.message,
            "evidence": issue.evidence,
        }
        for issue in issues
    ]

    selected_sequences = []

    for config in configs:
        selected_sequences.append({
            "sequence": config.sequence,
            "step_role": config.role,
            "set_order": config.set_order,
            "manual": bool(getattr(config, "manual", False)),
            "control": bool(getattr(config, "control", False)),
            "hardware_mode": normalize_hardware_mode(config.hardware_mode),
            "channel_count": get_sequence_batch_size(config),
            "channel_pattern": config.channel_pattern,
            "sequence_count": config.sequence_count,
            "volume_ul": config.volume_ul if config.role in ("aspirate", "dispense") else "",
            "liquid_class": config.liquid_class if config.role in ("aspirate", "dispense") else "",
            "transfer_type": config.transfer_type,
            "replicate": config.replicate,
            "marker": config.marker,
        })

    normalized_execution_plan = []

    for item in execution_plan:
        normalized_item = dict(item)
        normalized_item["action"] = humanize_type(normalized_item.get("action", ""))
        normalized_item["hardware_mode"] = normalize_hardware_mode(normalized_item.get("hardware_mode", "channels"))

        if "labware_format" in normalized_item:
            normalized_item["labware_format"] = humanize_type(normalized_item.get("labware_format", ""))

        normalized_execution_plan.append(normalized_item)

    normalized_layout_findings = []

    for item in layout_findings:
        if isinstance(item, dict):
            normalized_item = dict(item)
            if "type" in normalized_item:
                normalized_item["type"] = humanize_type(normalized_item["type"])
            normalized_layout_findings.append(normalized_item)
        else:
            normalized_layout_findings.append(item)

    normalized_trace_findings = []

    for item in trace_findings:
        if isinstance(item, dict):
            normalized_item = dict(item)
            if "type" in normalized_item:
                normalized_item["type"] = humanize_type(normalized_item["type"])
            normalized_trace_findings.append(normalized_item)
        else:
            normalized_trace_findings.append(item)

    return DevPalLiteOutput(
        lay_file=lay_file,
        layout_file=layout_file,
        trace_file=trace_file,
        selected_sequences=selected_sequences,
        source_positions=source_positions,
        destination_positions=destination_positions,
        transfer_pairs=[asdict(pair) for pair in pairs],
        execution_plan=normalized_execution_plan,
        marker_matches=marker_matches,
        layout_findings=normalized_layout_findings,
        trace_findings=normalized_trace_findings,
        validation_findings=validation_findings,
        issues=validation_findings,
        limitations=limitations,
        generated_review_script=review_script_path,
        report_path=report_path,
    )


def run_pyhamilton_simulation(
    lay_path: str,
    configs: List[SequenceStepConfig],
    execution_plan: List[Dict[str, Any]],
) -> Tuple[Optional[str], List[ValidationIssue]]:
    issues = []

    if phi is None or LayoutManager is None:
        issues.append(ValidationIssue(
            "critical",
            "runtime",
            f"pyhamilton is not available: {PYHAMILTON_IMPORT_ERROR}",
        ))
        return None, issues

    def get_plan_labware_name(
        item: Dict[str, Any],
        lay_metadata: Dict[str, Dict[str, Any]],
    ) -> str:
        for key in ("labware_name", "resolved_labware_name", "trace_or_method_labware", "deck_resource", "resource_name"):
            value = str(item.get(key) or "").strip()
            if value and value.lower() not in {"unknown", "none", "nan"}:
                return value

        seq = str(item.get("sequence") or "").strip()

        if seq in lay_metadata:
            value = str(lay_metadata[seq].get("labware_name") or "").strip()
            if value:
                return value

        for name, meta in lay_metadata.items():
            if str(name).strip().lower() == seq.lower():
                value = str(meta.get("labware_name") or "").strip()
                if value:
                    return value

        return seq

    def should_assign_labware_for_action(action: str, head_size: int) -> bool:
        return action in {"aspirate", "dispense"} and bool(head_size)

    run_start_time = time.time()
    lay_metadata = parse_lay_sequence_metadata(lay_path)
    resource_cache = {}

    try:
        layout_manager = LayoutManager(lay_path)

        with AutoRunHamiltonInterface(simulate=True) as ham_int:
            normal_logging(ham_int, os.getcwd())
            initialize(ham_int)

            for item in execution_plan:
                action = str(item.get("action") or "").strip()
                seq = str(item.get("sequence") or "").strip()
                mode = normalize_hardware_mode(item.get("hardware_mode") or "channels")
                head_size = get_head_size_from_mode(mode)

                try:
                    channel = normalize_assigned_channel_pattern(
                        item.get("channel_pattern") or "",
                        mode,
                    )
                except Exception:
                    channel = normalize_assigned_channel_pattern(
                        "111111111111",
                        "channels",
                    )

                volume = float(item.get("volume_ul") or 0)
                liquid_class = str(item.get("liquid_class") or "").strip()
                labware_name = get_plan_labware_name(item, lay_metadata)

                labware_resource = None

                if should_assign_labware_for_action(action, head_size):
                    resource_cache_key = str(labware_name or seq).strip().lower()

                    if resource_cache_key not in resource_cache:
                        resource_cache[resource_cache_key] = assign_sequence_labware_resource(
                            layout_manager,
                            labware_name or seq,
                        )

                    labware_resource = resource_cache[resource_cache_key]

                print(
                    f"RUNNING order={item.get('order')} "
                    f"set={item.get('set_order')} "
                    f"action={action} "
                    f"sequence={seq} "
                    f"hardware={mode} "
                    f"head_size={head_size or 'n/a'} "
                    f"labware={labware_name or 'unknown'} "
                    f"channel_count={channel.count('1') if channel else 0} "
                    f"liquid_class={liquid_class}"
                )

                if (
                    item.get("autoincrement")
                    and action in {"aspirate", "dispense"}
                    and not bool(item.get("manual"))
                ):
                    increment_amount = int(item.get("increment") or channel.count("1") or 1)
                    print(f"AUTO-INCREMENTING SEQUENCE {seq} BY {increment_amount}")
                    ph_inc_sequence(
                        ham_int,
                        sequence=seq,
                        increment=increment_amount,
                    )

                elif bool(item.get("manual")):
                    print(f"MANUAL SEQUENCE MODE: {seq} will not auto-increment.")

                if action == "tip_pick":
                    if head_size:
                        ph_tip_pick_up_head_seq(
                            ham_int,
                            tip_seq=seq,
                            head_size=head_size,
                        )
                    else:
                        ph_tip_pick_up_seq(
                            ham_int,
                            tipseq=seq,
                            channel=channel,
                        )

                elif action == "aspirate":
                    if head_size:
                        if labware_resource is None:
                            raise RuntimeError(
                                f"Could not assign head aspirate labware resource for sequence '{seq}' "
                                f"using labware '{labware_name}'."
                            )

                        ph_aspirate_head_seq(
                            ham_int,
                            plate=labware_resource,
                            head_asp_seq=seq,
                            vols=volume,
                            liq_class=liquid_class,
                            head_size=head_size,
                        )
                    else:
                        ph_aspirate_seq(
                            ham_int,
                            asp_seq=seq,
                            vols=volume,
                            channel=channel,
                            liq_class=liquid_class,
                        )

                elif action == "dispense":
                    if head_size:
                        if labware_resource is None:
                            raise RuntimeError(
                                f"Could not assign head dispense labware resource for sequence '{seq}' "
                                f"using labware '{labware_name}'."
                            )

                        ph_dispense_head_seq(
                            ham_int,
                            plate=labware_resource,
                            head_disp_seq=seq,
                            vols=volume,
                            liq_class=liquid_class,
                            head_size=head_size,
                        )
                    else:
                        ph_dispense_seq(
                            ham_int,
                            disp_seq=seq,
                            vols=volume,
                            channel=channel,
                            liq_class=liquid_class,
                        )

                elif action == "tip_eject":
                    if head_size:
                        ph_tip_eject_head(
                            ham_int,
                            head_size=head_size,
                        )
                    else:
                        ph_tip_eject_seq2(
                            ham_int,
                            waste_seq=seq or "Waste",
                            channel=channel,
                        )

    except (HamiltonStepError, HamiltonError, HamiltonTimeoutError, Exception) as exc:
        issues.append(ValidationIssue(
            "critical",
            "runtime",
            f"PyHamilton simulation failed: {type(exc).__name__}: {exc}",
        ))

    time.sleep(5)

    trace_path = get_latest_trc_file(
        run_start_time=run_start_time,
        wait_seconds=45,
        trace_dir=DEFAULT_TRACE_DIR,
        require_after_start=False,
    )

    if not trace_path:
        issues.append(ValidationIssue(
            "major",
            "trace",
            "No STAR_OEM_noFan trace file found after PyHamilton simulation.",
            {"trace_dir": DEFAULT_TRACE_DIR},
        ))
    else:
        issues.append(ValidationIssue(
            "info",
            "trace",
            "Using latest STAR_OEM_noFan trace generated/found after PyHamilton simulation.",
            {
                "trace_path": trace_path,
                "run_start_time": run_start_time,
                "trace_modified_time": os.path.getmtime(trace_path),
            },
        ))

    return trace_path, issues


class DevPalLiteApp:
    def __init__(self):
        self.root = tk.Tk()
        self.root.title("DevPal Lite Sequence Layout Runner")
        self.root.geometry("1180x780")
        self.root.minsize(1000, 650)
        self.root.option_add("*Font", APP_FONT)

        self.lay_path = ""
        self.layout_path = ""
        self.trace_path = ""

        self.sequence_names = []
        self.sequence_counts = {}
        self.liquid_classes = []

        self.sequence_vars = []
        self.sequence_role_vars = {}
        self.sequence_channel_vars = {}
        self.sequence_hardware_vars = {}
        self.sequence_count_vars = {}
        self.sequence_manual_vars = {}
        self.sequence_control_vars = {}
        self.sequence_set_vars = {}
        self.sequence_set_widgets = {}

        self.liquid_class_var = tk.StringVar()
        self.transfer_type_var = tk.StringVar(value="samples")
        self.replicate_var = tk.StringVar(value="single")
        self.volume_var = tk.StringVar(value="100")
        self.marker_var = tk.StringVar()
        self.report_format_var = tk.StringVar(value="docx")

        self.tracking_var = tk.IntVar(value=0)
        self.variant_var = tk.IntVar(value=0)

        self.set_header_widget = None
        self.set_liquid_classes = []
        self.set_liquid_listbox = None
        self.set_liquid_summary_var = tk.StringVar(value="No set liquid classes selected.")

        self.background_image_original = None
        self.background_image_tk = None

        self._build_ui()

    def style_label(self, widget):
        widget.configure(
            bg=self.UI_TEXT_BG,
            fg=self.UI_TEXT_FG,
            relief="raised",
            bd=2,
            padx=4,
            pady=2,
        )
        return widget

    def style_plain_label(self, widget):
        widget.configure(
            bg=self.UI_BG,
            fg=self.UI_TEXT_FG,
            relief="flat",
            bd=0,
            padx=2,
            pady=1,
        )
        return widget

    def style_frame(self, widget):
        widget.configure(bg=self.UI_BG, relief="ridge", bd=2)
        return widget

    def style_button(self, widget):
        widget.configure(
            bg=self.UI_BUTTON_BG,
            fg=self.UI_TEXT_FG,
            activebackground=self.UI_BUTTON_ACTIVE,
            activeforeground=self.UI_TEXT_FG,
            relief="raised",
            bd=3,
            padx=5,
            pady=2,
        )
        return widget

    def make_bg_label(self, parent, text, **kwargs):
        label = tk.Label(
            parent,
            text=text,
            bg=self.UI_TEXT_BG,
            fg=self.UI_TEXT_FG,
            relief="raised",
            bd=2,
            padx=4,
            pady=2,
            **kwargs,
        )
        return label

    def refresh_checkbutton_state(self, checkbox, variable):
        if variable.get():
            checkbox.configure(selectcolor="#1e6bff")
        else:
            checkbox.configure(selectcolor="white")

    def make_checkbutton(self, parent, text="", variable=None, command=None, borderless=False):
        checkbox = tk.Checkbutton(
            parent,
            text=text,
            variable=variable,
            bg=UI_BG,
            fg=UI_TEXT_FG,
            activebackground=UI_BUTTON_ACTIVE,
            activeforeground=UI_TEXT_FG,
            selectcolor="white",
            indicatoron=True,
            highlightthickness=0,
            anchor="w",
            padx=2,
            pady=1,
        )

        if borderless:
            checkbox.configure(bd=0, relief="flat")
        else:
            checkbox.configure(bd=2, relief="raised")

        def wrapped_command():
            self.refresh_checkbutton_state(checkbox, variable)
            if command:
                command()

        checkbox.configure(command=wrapped_command)
        self.refresh_checkbutton_state(checkbox, variable)
        return checkbox

    def resize_background_image(self, event=None):
        if self.background_image_original is None or ImageTk is None:
            return

        try:
            width = max(self.background_canvas.winfo_width(), 1)
            height = max(self.background_canvas.winfo_height(), 1)

            image = self.background_image_original.copy()
            image_ratio = image.width / image.height
            canvas_ratio = width / height

            if image_ratio > canvas_ratio:
                new_height = height
                new_width = int(height * image_ratio)
            else:
                new_width = width
                new_height = int(width / image_ratio)

            image = image.resize((new_width, new_height))
            left = max((new_width - width) // 2, 0)
            top = max((new_height - height) // 2, 0)
            image = image.crop((left, top, left + width, top + height))

            self.background_image_tk = ImageTk.PhotoImage(image)
            self.background_canvas.delete("all")
            self.background_canvas.create_image(0, 0, image=self.background_image_tk, anchor="nw")

        except Exception:
            pass

    def load_right_background_image(self):
        self.right_bg_original = None
        self.right_bg_photo = None

        image_path = os.path.join(os.getcwd(), "backgroundSA.png")

        if not os.path.exists(image_path):
            return

        if Image is None or ImageTk is None:
            return

        try:
            self.right_bg_original = Image.open(image_path).convert("RGB")
        except Exception:
            self.right_bg_original = None

    def resize_right_background(self, event=None):
        if self.right_bg_original is None:
            return

        width = max(self.right_bg_frame.winfo_width(), 1)
        height = max(self.right_bg_frame.winfo_height(), 1)

        if width <= 1 or height <= 1:
            return

        image = self.right_bg_original.copy()
        image_ratio = image.width / image.height
        frame_ratio = width / height

        if frame_ratio > image_ratio:
            new_width = width
            new_height = int(width / image_ratio)
        else:
            new_height = height
            new_width = int(height * image_ratio)

        image = image.resize((new_width, new_height), Image.LANCZOS)

        left = max((new_width - width) // 2, 0)
        top = max((new_height - height) // 2, 0)
        image = image.crop((left, top, left + width, top + height))

        self.right_bg_photo = ImageTk.PhotoImage(image)
        self.right_bg_label.configure(image=self.right_bg_photo)

    def _build_ui(self):
        self.root.configure(bg=UI_BG)

        icon_path = os.path.join(os.getcwd(), "thumbs_up.ico")
        if os.path.exists(icon_path):
            try:
                self.root.iconbitmap(icon_path)
            except Exception:
                pass

        self.main_container = tk.Frame(self.root, bg=UI_BG)
        self.main_container.pack(fill="both", expand=True)

        self.main_container.grid_rowconfigure(0, weight=1)
        self.main_container.grid_columnconfigure(0, weight=3)
        self.main_container.grid_columnconfigure(1, weight=2)

        self.left_ui_frame = tk.Frame(self.main_container, bg=UI_BG)
        self.left_ui_frame.grid(row=0, column=0, sticky="nsew")

        self.right_image_frame = tk.Frame(self.main_container, bg=UI_BG)
        self.right_image_frame.grid(row=0, column=1, sticky="nsew")

        self.background_canvas = tk.Canvas(
            self.right_image_frame,
            bg=UI_BG,
            highlightthickness=0,
            bd=0,
        )
        self.background_canvas.pack(fill="both", expand=True)

        background_path = os.path.join(os.getcwd(), "backgroundSA.png")
        if os.path.exists(background_path) and Image is not None:
            try:
                self.background_image_original = Image.open(background_path)
            except Exception:
                self.background_image_original = None

        self.background_canvas.bind("<Configure>", self.resize_background_image)

        def purple_label(parent, text, **kwargs):
            return tk.Label(
                parent,
                text=text,
                bg=UI_TEXT_BG,
                fg=UI_TEXT_FG,
                bd=2,
                relief="raised",
                padx=5,
                pady=2,
                **kwargs,
            )

        def purple_button(parent, text, command, **kwargs):
            return tk.Button(
                parent,
                text=text,
                command=command,
                bg=UI_BUTTON_BG,
                fg=UI_TEXT_FG,
                activebackground=UI_BUTTON_ACTIVE,
                activeforeground=UI_TEXT_FG,
                bd=3,
                relief="raised",
                padx=8,
                pady=3,
                **kwargs,
            )

        top = tk.Frame(self.left_ui_frame, bg=UI_BG)
        top.pack(fill="x", padx=8, pady=(8, 4))

        lay_row = tk.Frame(top, bg=UI_BG)
        lay_row.pack(fill="x", pady=2)

        purple_button(lay_row, "Upload .lay File", self.upload_lay, width=18).pack(side="left", padx=(0, 6))
        self.lay_label = tk.Label(lay_row, text="No .lay selected", bg=UI_BG, fg=UI_TEXT_FG, anchor="w", bd=2,
                                  relief="sunken", padx=5)
        self.lay_label.pack(side="left", fill="x", expand=True)

        layout_row = tk.Frame(top, bg=UI_BG)
        layout_row.pack(fill="x", pady=2)

        purple_button(layout_row, "Upload Plate Layout", self.upload_layout, width=18).pack(side="left", padx=(0, 6))
        self.layout_label = tk.Label(layout_row, text="No layout selected", bg=UI_BG, fg=UI_TEXT_FG, anchor="w", bd=2,
                                     relief="sunken", padx=5)
        self.layout_label.pack(side="left", fill="x", expand=True)

        form = tk.LabelFrame(self.left_ui_frame, text="Required Transfer Information", bg=UI_BG, fg=UI_TEXT_FG, bd=3,
                             relief="ridge")
        form.pack(fill="x", padx=8, pady=4)

        form_inner = tk.Frame(form, bg=UI_BG)
        form_inner.pack(fill="x", padx=6, pady=6)

        purple_label(form_inner, "Transfer Type").grid(row=0, column=0, sticky="w", padx=(0, 4), pady=3)
        ttk.Combobox(form_inner, textvariable=self.transfer_type_var, values=TRANSFER_OPTIONS, state="readonly",
                     width=12).grid(row=0, column=1, sticky="w", padx=(0, 8), pady=3)

        purple_label(form_inner, "Replicate").grid(row=0, column=2, sticky="w", padx=(0, 4), pady=3)
        ttk.Combobox(form_inner, textvariable=self.replicate_var, values=REPLICATE_OPTIONS, state="readonly",
                     width=10).grid(row=0, column=3, sticky="w", padx=(0, 8), pady=3)

        purple_label(form_inner, "Volume uL").grid(row=0, column=4, sticky="w", padx=(0, 4), pady=3)
        tk.Entry(form_inner, textvariable=self.volume_var, width=9, bg="white", fg="black", bd=2, relief="sunken").grid(
            row=0, column=5, sticky="w", padx=(0, 8), pady=3)

        purple_label(form_inner, "Marker").grid(row=0, column=6, sticky="w", padx=(0, 4), pady=3)
        tk.Entry(form_inner, textvariable=self.marker_var, width=9, bg="white", fg="black", bd=2, relief="sunken").grid(
            row=0, column=7, sticky="w", padx=(0, 0), pady=3)

        lc_row = tk.Frame(form, bg=UI_BG)
        lc_row.pack(fill="x", padx=6, pady=(0, 6))

        purple_label(lc_row, "Liquid Class").pack(side="left", padx=(0, 6))

        self.liquid_combo = SearchableCombobox(lc_row, textvariable=self.liquid_class_var, completevalues=[], width=50)
        self.liquid_combo.pack(side="left", padx=(0, 6))

        purple_button(lc_row, "Add LC", self.add_liquid_class_for_next_set, width=8).pack(side="left", padx=(0, 4))
        purple_button(lc_row, "Clear", self.clear_set_liquid_classes, width=7).pack(side="left")

        self.set_liquid_summary_box = tk.LabelFrame(self.left_ui_frame, text="Liquid Class Set Assignments", bg=UI_BG,
                                                    fg=UI_TEXT_FG, bd=3, relief="ridge")
        self.set_liquid_summary_box.pack(fill="x", padx=8, pady=4)

        self.set_liquid_summary_label = tk.Label(self.set_liquid_summary_box, textvariable=self.set_liquid_summary_var,
                                                 bg=UI_BG, fg=UI_TEXT_FG, anchor="w", padx=6, pady=3)
        self.set_liquid_summary_label.pack(fill="x", padx=4, pady=(3, 2))

        self.set_liquid_listbox = tk.Listbox(self.set_liquid_summary_box, height=3, width=70, bg=UI_BG, fg=UI_TEXT_FG,
                                             selectbackground=UI_BUTTON_ACTIVE, selectforeground=UI_TEXT_FG, bd=2,
                                             relief="sunken")
        self.set_liquid_listbox.pack(fill="x", padx=4, pady=(0, 5))

        options_row = tk.Frame(self.left_ui_frame, bg=UI_BG)
        options_row.pack(fill="x", padx=8, pady=4)

        self.make_checkbutton(options_row, text="Tracking Enabled", variable=self.tracking_var).pack(side="left",
                                                                                                     padx=(0, 8))
        self.make_checkbutton(options_row, text="Variant", variable=self.variant_var,
                              command=self.update_set_column_visibility).pack(side="left", padx=(0, 12))

        report_box = tk.LabelFrame(options_row, text="Report Options", bg=UI_BG, fg=UI_TEXT_FG, bd=3, relief="ridge")
        report_box.pack(side="left")

        purple_label(report_box, "Report Format").pack(side="left", padx=5, pady=5)
        ttk.Combobox(report_box, textvariable=self.report_format_var, values=["docx", "pdf", "txt"], state="readonly",
                     width=8).pack(side="left", padx=(0, 6), pady=5)

        list_frame = tk.LabelFrame(self.left_ui_frame, text="Extracted Method/Deck Sequences", bg=UI_BG, fg=UI_TEXT_FG,
                                   bd=3, relief="ridge")
        list_frame.pack(fill="both", expand=True, padx=8, pady=4)

        self.sequence_canvas = tk.Canvas(list_frame, bg=UI_BG, highlightthickness=0, bd=0)
        y_scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=self.sequence_canvas.yview)
        x_scrollbar = ttk.Scrollbar(list_frame, orient="horizontal", command=self.sequence_canvas.xview)

        self.sequence_frame = tk.Frame(self.sequence_canvas, bg=UI_BG)
        self.scroll_frame = self.sequence_frame
        self.sequence_window = self.sequence_canvas.create_window((0, 0), window=self.sequence_frame, anchor="nw")

        self.sequence_frame.bind("<Configure>", lambda event: self.sequence_canvas.configure(
            scrollregion=self.sequence_canvas.bbox("all")))

        self.sequence_canvas.configure(yscrollcommand=y_scrollbar.set, xscrollcommand=x_scrollbar.set)

        self.sequence_canvas.grid(row=0, column=0, sticky="nsew")
        y_scrollbar.grid(row=0, column=1, sticky="ns")
        x_scrollbar.grid(row=1, column=0, sticky="ew")

        list_frame.grid_rowconfigure(0, weight=1)
        list_frame.grid_columnconfigure(0, weight=1)

        bottom = tk.Frame(self.left_ui_frame, bg=UI_BG)
        bottom.pack(fill="x", padx=8, pady=(4, 8))

        purple_button(bottom, "Run Sequence Analyzer", self.generate_plan_run_and_validate, width=24).pack(side="left",
                                                                                                           padx=(0, 6))
        purple_button(bottom, "Select Existing .trc", self.upload_trace, width=17).pack(side="left", padx=(0, 6))

        self.status = tk.Label(bottom, text="Upload .lay file to begin.", bg=UI_BG, fg=UI_TEXT_FG, anchor="w", bd=2,
                               relief="sunken", padx=5)
        self.status.pack(side="left", fill="x", expand=True)

    def _load_liquid_classes(self):
        return self.load_liquid_classes()

    def load_liquid_classes(self):
        self.liquid_classes = load_combined_liquid_classes(self.lay_path)

        self.liquid_combo.completevalues = self.liquid_classes
        self.liquid_combo["values"] = self.liquid_classes

        current_value = self.liquid_class_var.get().strip()

        if current_value not in self.liquid_classes:
            self.liquid_class_var.set("")

        if self.liquid_classes:
            self.liquid_class_var.set(self.liquid_classes[0])
        else:
            messagebox.showwarning(
                "Liquid Class Warning",
                "No liquid classes were found in the uploaded method files. "
                "DevPal Lite will not use default LC.mdb/LC.accdb values.",
            )

    def add_liquid_class_for_next_set(self):
        liquid_class = self.liquid_class_var.get().strip()

        if not liquid_class:
            messagebox.showerror("Missing Liquid Class", "Select a liquid class first.")
            return

        if liquid_class not in self.liquid_classes:
            messagebox.showerror(
                "Invalid Liquid Class",
                f"Selected liquid class is not from the uploaded method:\n\n{liquid_class}",
            )
            return

        self.set_liquid_classes.append(liquid_class)
        self.refresh_set_liquid_class_display()

    def clear_set_liquid_classes(self):
        self.set_liquid_classes = []
        self.refresh_set_liquid_class_display()

    def refresh_set_liquid_class_display(self):
        if self.set_liquid_listbox is not None:
            self.set_liquid_listbox.delete(0, tk.END)

            for index, liquid_class in enumerate(self.set_liquid_classes, start=1):
                self.set_liquid_listbox.insert(tk.END, f"Set {index}: {liquid_class}")

        if self.set_liquid_classes:
            self.set_liquid_summary_var.set(
                "Liquid classes assigned by click order: "
                + "; ".join(
                    f"Set {index} = {liquid_class}"
                    for index, liquid_class in enumerate(self.set_liquid_classes, start=1)
                )
            )
        else:
            self.set_liquid_summary_var.set("No set liquid classes selected.")

    def upload_lay(self):
        path = filedialog.askopenfilename(
            title="Select .lay File",
            filetypes=[("LAY files", "*.lay"), ("All files", "*.*")],
        )

        if not path:
            return

        self.lay_path = path
        self.lay_label.config(text=path)

        try:
            self.sequence_names = sorted(parse_lay_file(path), key=lambda x: str(x).lower())
            self.sequence_counts = parse_lay_sequence_counts(path)

            if not self.sequence_names:
                messagebox.showerror(
                    "LAY Parse Error",
                    "No valid sequences were found in the selected .lay file.",
                )
                return

            self.load_liquid_classes()
            self.render_sequences()

            self.status.config(
                text=f"Loaded {len(self.sequence_names)} sequences and {len(self.liquid_classes)} liquid classes.",
            )

        except Exception as exc:
            messagebox.showerror("LAY Parse Error", str(exc))

    def upload_layout(self):
        path = filedialog.askopenfilename(
            title="Select Plate Layout File",
            filetypes=[("Plate layout files", "*.csv *.xls *.xlsx *.docx"), ("All files", "*.*")],
        )

        if path:
            self.layout_path = path
            self.layout_label.config(text=path)

    def upload_trace(self):
        path = filedialog.askopenfilename(
            title="Select Trace File",
            filetypes=[("Trace files", "*.trc"), ("All files", "*.*")],
        )

        if path:
            self.trace_path = path
            self.status.config(text=f"Selected trace: {path}")

    def get_sequence_frame(self):
        if hasattr(self, "sequence_frame") and self.sequence_frame is not None:
            return self.sequence_frame

        if hasattr(self, "sequence_table_frame") and self.sequence_table_frame is not None:
            self.sequence_frame = self.sequence_table_frame
            return self.sequence_frame

        if hasattr(self, "scrollable_frame") and self.scrollable_frame is not None:
            self.sequence_frame = self.scrollable_frame
            return self.sequence_frame

        raise AttributeError(
            "No sequence frame exists. Create self.sequence_frame in build_ui() before calling render_sequences().")

    def render_sequences(self):
        if not hasattr(self, "sequence_frame") or self.sequence_frame is None:
            messagebox.showerror("UI Error", "Sequence frame does not exist. Rebuild UI before rendering sequences.")
            return

        sequence_frame = self.sequence_frame

        for widget in sequence_frame.winfo_children():
            widget.destroy()

        self.sequence_vars = []
        self.sequence_role_vars = {}
        self.sequence_channel_vars = {}
        self.sequence_hardware_vars = {}
        self.sequence_count_vars = {}
        self.sequence_manual_vars = {}
        self.sequence_control_vars = {}
        self.sequence_set_vars = {}
        self.sequence_set_widgets = {}
        self.set_header_widget = None

        headers = [
            ("Use", 0, 5),
            ("Control", 1, 7),
            ("Set", 2, 4),
            ("Sequence", 3, 48),
            ("Step Role", 4, 12),
            ("Manual", 5, 7),
            ("Channel Pattern", 6, 16),
            ("Hardware", 7, 10),
            ("Seq Count", 8, 9),
        ]

        for text, col, width in headers:
            label = tk.Label(
                sequence_frame,
                text=text,
                width=width,
                anchor="center",
                bg=UI_TEXT_BG,
                fg=UI_TEXT_FG,
                font=("Arial", 9, "bold"),
                bd=2,
                relief="raised",
            )
            label.grid(row=0, column=col, padx=3, pady=4, sticky="w")

            if text == "Set":
                self.set_header_widget = label

        sequence_names = list(getattr(self, "sequence_names", []) or [])
        sequence_counts = getattr(self, "sequence_counts", {}) or {}

        longest_name_len = max([len(str(name)) for name in sequence_names], default=30)
        sequence_width = max(48, min(longest_name_len + 6, 90))

        for row_index, sequence_name in enumerate(sequence_names, start=1):
            sequence_count = int(sequence_counts.get(sequence_name, 0) or 0)

            default_hardware = detect_hardware_mode(sequence_name, "").lower()
            if default_hardware not in ("channels", "head"):
                default_hardware = "channels"

            default_channel_pattern = "1" * 96 if default_hardware == "head" else "111111111111"

            use_var = tk.IntVar(value=0)
            control_var = tk.IntVar(value=0)
            set_var = tk.StringVar(value="")
            role_var = tk.StringVar(value="aspirate")
            manual_var = tk.IntVar(value=0)
            channel_var = tk.StringVar(value=default_channel_pattern)
            hardware_var = tk.StringVar(value=default_hardware)
            sequence_count_var = tk.StringVar(value=str(sequence_count))

            self.sequence_vars.append((sequence_name, use_var))
            self.sequence_control_vars[sequence_name] = control_var
            self.sequence_role_vars[sequence_name] = role_var
            self.sequence_manual_vars[sequence_name] = manual_var
            self.sequence_channel_vars[sequence_name] = channel_var
            self.sequence_hardware_vars[sequence_name] = hardware_var
            self.sequence_count_vars[sequence_name] = sequence_count_var
            self.sequence_set_vars[sequence_name] = set_var

            self.make_checkbutton(
                sequence_frame,
                variable=use_var,
                command=self.update_set_column_visibility,
                borderless=True,
            ).grid(row=row_index, column=0, padx=3, pady=1, sticky="w")

            self.make_checkbutton(
                sequence_frame,
                variable=control_var,
                borderless=True,
            ).grid(row=row_index, column=1, padx=3, pady=1, sticky="w")

            set_entry = tk.Entry(
                sequence_frame,
                textvariable=set_var,
                width=4,
                bg="white",
                fg="black",
                bd=1,
                relief="sunken",
                justify="center",
            )
            set_entry.grid(row=row_index, column=2, padx=3, pady=1, sticky="w")
            self.sequence_set_widgets[sequence_name] = set_entry

            tk.Label(
                sequence_frame,
                text=sequence_name,
                width=sequence_width,
                anchor="w",
                bg=UI_BG,
                fg=UI_TEXT_FG,
            ).grid(row=row_index, column=3, padx=3, pady=1, sticky="w")

            role_combo = ttk.Combobox(
                sequence_frame,
                textvariable=role_var,
                values=ROLE_OPTIONS,
                width=12,
                state="readonly",
            )
            role_combo.grid(row=row_index, column=4, padx=3, pady=1, sticky="w")

            self.make_checkbutton(
                sequence_frame,
                variable=manual_var,
                borderless=True,
            ).grid(row=row_index, column=5, padx=3, pady=1, sticky="w")

            tk.Entry(
                sequence_frame,
                textvariable=channel_var,
                width=16,
                bg="white",
                fg="black",
                bd=1,
                relief="sunken",
            ).grid(row=row_index, column=6, padx=3, pady=1, sticky="w")

            hardware_combo = ttk.Combobox(
                sequence_frame,
                textvariable=hardware_var,
                values=["channels", "head"],
                width=10,
                state="readonly",
            )
            hardware_combo.grid(row=row_index, column=7, padx=3, pady=1, sticky="w")

            tk.Label(
                sequence_frame,
                textvariable=sequence_count_var,
                width=9,
                anchor="center",
                bg=UI_BG,
                fg=UI_TEXT_FG,
            ).grid(row=row_index, column=8, padx=3, pady=1, sticky="w")

            def on_role_or_hardware_change(event=None, seq=sequence_name):
                role = self.sequence_role_vars[seq].get().strip()
                hardware = self.sequence_hardware_vars[seq].get().strip().lower()

                detected = detect_hardware_mode(seq, role).lower()
                if detected in ("channels", "head"):
                    hardware = detected
                    self.sequence_hardware_vars[seq].set(hardware)

                current_pattern = self.sequence_channel_vars[seq].get().strip()

                if hardware == "head":
                    if not current_pattern or len(current_pattern) != 96:
                        self.sequence_channel_vars[seq].set("1" * 96)
                else:
                    if not current_pattern or len(current_pattern) == 96:
                        self.sequence_channel_vars[seq].set("111111111111")

            role_combo.bind("<<ComboboxSelected>>", on_role_or_hardware_change)
            hardware_combo.bind("<<ComboboxSelected>>", on_role_or_hardware_change)

        if not sequence_names:
            tk.Label(
                sequence_frame,
                text="No sequences loaded.",
                bg=UI_BG,
                fg=UI_TEXT_FG,
                anchor="w",
            ).grid(row=1, column=0, columnspan=9, padx=4, pady=8, sticky="w")

        self.update_set_column_visibility()

        try:
            sequence_frame.update_idletasks()
            self.sequence_canvas.configure(scrollregion=self.sequence_canvas.bbox("all"))
        except Exception:
            pass

    def update_set_column_visibility(self):
        selected_count = sum(1 for _, use_var in self.sequence_vars if use_var.get())
        show_sets = selected_count > 3 or bool(self.variant_var.get())

        if self.set_header_widget is not None:
            if show_sets:
                self.set_header_widget.grid()
            else:
                self.set_header_widget.grid_remove()

        for sequence_name, widget in self.sequence_set_widgets.items():
            if show_sets:
                widget.grid()
            else:
                widget.grid_remove()
                if sequence_name in self.sequence_set_vars:
                    self.sequence_set_vars[sequence_name].set("")

        try:
            if hasattr(self, "sequence_canvas"):
                self.sequence_canvas.configure(scrollregion=self.sequence_canvas.bbox("all"))
        except Exception:
            pass

    def collect_configs(self) -> Optional[List[SequenceStepConfig]]:
        configs = []

        marker = self.marker_var.get().strip().upper()
        transfer_type = self.transfer_type_var.get().strip() or "samples"
        replicate = self.replicate_var.get().strip().lower()
        variant_mode = bool(self.variant_var.get())

        selected_sequences = [
            sequence_name
            for sequence_name, use_var in self.sequence_vars
            if use_var.get()
        ]

        if not selected_sequences:
            return []

        use_set_order = len(selected_sequences) > 3 or variant_mode

        raw_volume_text = self.volume_var.get().strip()

        if not raw_volume_text:
            messagebox.showerror("Invalid Volume", "Volume is required.")
            return None

        try:
            volume_values = [
                float(value.strip())
                for value in raw_volume_text.split(",")
                if value.strip()
            ]
        except Exception:
            messagebox.showerror(
                "Invalid Volume",
                "Volume must be numeric. For multiple sets, separate volumes with commas, such as 100, 200.",
            )
            return None

        if not volume_values:
            messagebox.showerror("Invalid Volume", "Volume is required.")
            return None

        for volume in volume_values:
            if volume <= 0 or volume > 1000:
                messagebox.showerror(
                    "Invalid Volume",
                    f"Volume must be between 1 and 1000 uL. Invalid value: {volume}",
                )
                return None

        selected_set_numbers = []

        if use_set_order:
            for sequence_name in selected_sequences:
                raw_set = self.sequence_set_vars[sequence_name].get().strip()

                if not raw_set:
                    messagebox.showerror(
                        "Missing Set Number",
                        f"Set number is required for selected sequence: {sequence_name}",
                    )
                    return None

                try:
                    set_num = int(raw_set)
                except Exception:
                    messagebox.showerror(
                        "Invalid Set Number",
                        f"Set number must be numeric for sequence: {sequence_name}",
                    )
                    return None

                if set_num <= 0:
                    messagebox.showerror(
                        "Invalid Set Number",
                        f"Set number must be 1 or higher for sequence: {sequence_name}",
                    )
                    return None

                selected_set_numbers.append(set_num)

        required_set_count = max(selected_set_numbers) if selected_set_numbers else 1

        if len(volume_values) not in {1, required_set_count}:
            messagebox.showerror(
                "Volume / Set Mismatch",
                (
                    f"You entered {len(volume_values)} volume value(s), but the selected sequences use "
                    f"{required_set_count} set(s).\n\n"
                    "Enter one volume to use for all sets, or enter one comma-separated volume per set."
                ),
            )
            return None

        if required_set_count > 1 and len(self.set_liquid_classes) < required_set_count:
            messagebox.showerror(
                "Missing Set Liquid Classes",
                (
                    f"You have {required_set_count} set(s), but only "
                    f"{len(self.set_liquid_classes)} liquid class assignment(s).\n\n"
                    "Select each liquid class in order and click 'Add LC to Next Set'."
                ),
            )
            return None

        lay_metadata = parse_lay_sequence_metadata(self.lay_path) if self.lay_path else {}

        for selected_index, sequence_name in enumerate(selected_sequences):
            role = self.sequence_role_vars[sequence_name].get().strip()
            raw_channel_pattern = self.sequence_channel_vars[sequence_name].get().strip()
            hardware_mode = normalize_hardware_mode(self.sequence_hardware_vars[sequence_name].get())
            manual = bool(self.sequence_manual_vars[sequence_name].get())

            if hasattr(self, "sequence_control_vars") and sequence_name in self.sequence_control_vars:
                control = bool(self.sequence_control_vars[sequence_name].get())
            else:
                control = False

            try:
                sequence_count = int(self.sequence_count_vars[sequence_name].get())
            except Exception:
                sequence_count = int(self.sequence_counts.get(sequence_name, 0) or 0)

            set_order = 1

            if use_set_order:
                set_order = int(self.sequence_set_vars[sequence_name].get().strip())

            if len(volume_values) == 1:
                volume_ul = volume_values[0]
            else:
                volume_ul = volume_values[set_order - 1]

            if role in ("aspirate", "dispense"):
                if self.set_liquid_classes:
                    if set_order <= len(self.set_liquid_classes):
                        liquid_class = self.set_liquid_classes[set_order - 1]
                    else:
                        messagebox.showerror(
                            "Missing Set Liquid Class",
                            f"No liquid class assigned for Set {set_order}.",
                        )
                        return None
                else:
                    liquid_class = self.liquid_class_var.get().strip()

                if not liquid_class:
                    messagebox.showerror(
                        "Missing Liquid Class",
                        "Liquid class is required for aspirate/dispense steps.",
                    )
                    return None

                if liquid_class not in self.liquid_classes:
                    messagebox.showerror(
                        "Invalid Liquid Class",
                        f"Liquid class is not from the uploaded method:\n\n{liquid_class}",
                    )
                    return None
            else:
                liquid_class = ""

            try:
                if hardware_mode in {"head96", "head384"}:
                    channel_pattern = normalize_assigned_channel_pattern("", hardware_mode)
                else:
                    channel_pattern = normalize_assigned_channel_pattern(raw_channel_pattern, "channels")
            except Exception as exc:
                messagebox.showerror(
                    "Invalid Channel Pattern",
                    f"{sequence_name} has invalid channel pattern.\n\n{exc}",
                )
                return None

            configs.append(
                SequenceStepConfig(
                    sequence=sequence_name,
                    role=role,
                    liquid_class=liquid_class,
                    transfer_type=transfer_type,
                    replicate=replicate,
                    volume_ul=volume_ul,
                    marker=marker,
                    channel_pattern=channel_pattern,
                    hardware_mode=hardware_mode,
                    selected_index=selected_index,
                    sequence_count=sequence_count,
                    manual=manual,
                    set_order=set_order,
                    control=control,
                )
            )

        validation_error = validate_step_sets(configs, variant_mode, use_set_order)

        if validation_error:
            messagebox.showerror("Invalid Step Sets", validation_error)
            return None

        tip_issues = validate_tip_volume_rules(configs, lay_metadata)

        if tip_issues:
            messagebox.showerror(
                "Tip Type / Volume Rule Error",
                "\n".join(unique_issue_messages(tip_issues)[:10]),
            )
            return None

        return configs

    def generate_plan_run_and_validate(self):
        if not self.lay_path:
            messagebox.showerror("Missing LAY", "Upload a .lay file first.")
            return

        if not self.layout_path:
            messagebox.showerror("Missing Plate Layout", "Upload a plate layout file first.")
            return

        configs = self.collect_configs()

        if configs is None:
            return

        if not configs:
            messagebox.showerror("Missing Sequences", "Select sequences first.")
            return

        issues = validate_step_configs(configs)
        lay_metadata = parse_lay_sequence_metadata(self.lay_path)

        issues.extend(validate_tip_volume_rules(
            configs,
            lay_metadata,
        ))
        critical = [issue for issue in issues if issue.severity == "critical"]

        if critical:
            messagebox.showerror(
                "Configuration Error",
                "\n".join(issue.message for issue in critical[:10]),
            )
            return

        today_folder = datetime.now().strftime("%d%b%Y")
        output_dir = os.path.join(DEFAULT_RESULTS_DIR, today_folder)
        os.makedirs(output_dir, exist_ok=True)

        marker = configs[0].marker if configs else ""
        marker_matches = {}
        layout_findings = []
        trace_findings = []
        limitations = []

        execution_plan = build_execution_plan(
            configs=configs,
            lay_metadata=lay_metadata,
        )

        try:
            marker_matches = parse_plate_layout_file(self.layout_path, marker)
        except Exception as exc:
            issues.append(ValidationIssue(
                "critical",
                "layout",
                f"Could not parse plate layout: {exc}",
            ))

        pairs = generate_deterministic_transfer_pairs(configs, marker_matches)

        if marker_matches:
            layout_findings, layout_issues = validate_layout_matches(
                marker_matches,
                pairs,
                marker,
            )
            issues.extend(layout_issues)
        else:
            limitations.append("No marker matches were found in the plate layout.")

        labware_summary = build_labware_validation_summary(
            marker_matches,
            trace_findings,
            configs,
            lay_metadata,
        )

        trace_findings.extend(labware_summary)

        for labware_item in labware_summary:
            if labware_item.get("status") == "fail":
                issues.append(ValidationIssue(
                    "critical",
                    "URGENT_LABWARE_MISMATCH",
                    labware_item.get("message", "Wrong labware detected."),
                    labware_item,
                ))

        review_script = generate_pyhamilton_review_script(
            configs,
            execution_plan,
            output_dir,
        )

        self.status.config(text="Running PyHamilton simulation...")
        self.root.update_idletasks()

        uploaded_trace_before_run = self.trace_path

        runtime_trace, runtime_issues = run_pyhamilton_simulation(
            self.lay_path,
            configs,
            execution_plan,
        )

        issues.extend(runtime_issues)

        if runtime_trace:
            self.trace_path = runtime_trace
        elif uploaded_trace_before_run:
            self.trace_path = uploaded_trace_before_run
            issues.append(ValidationIssue(
                "major",
                "trace",
                "No new simulation trace was found, so DevPal Lite used the manually uploaded trace as fallback.",
                {"uploaded_trace": uploaded_trace_before_run},
            ))
        else:
            self.trace_path = ""
            issues.append(ValidationIssue(
                "major",
                "trace",
                "No new STAR_OEM_noFan trace was found after PyHamilton simulation.",
                {"trace_directory": DEFAULT_TRACE_DIR},
            ))

        if self.trace_path:
            try:
                aspirate_steps, dispense_steps, ordered_steps = parse_trc_file(
                    self.trace_path
                )

                trace_text = read_text_file_best_effort(self.trace_path)

                runtime_step_findings, runtime_step_issues = validate_expected_runtime_steps(
                    execution_plan,
                    ordered_steps,
                    trace_text,
                )

                trace_findings.extend(runtime_step_findings)
                issues.extend(runtime_step_issues)

                failed_runtime_steps = [
                    issue for issue in runtime_step_issues
                    if issue.severity in ("critical", "major")
                ]

                if failed_runtime_steps:
                    messagebox.showerror(
                        "Runtime Step Failure",
                        "\n".join(issue.message for issue in failed_runtime_steps[:3]),
                    )

                ok, cross_match = cross_match_positions(
                    aspirate_steps,
                    dispense_steps,
                    configs[0].replicate,
                )

                trace_findings.append({
                    "type": "trace_summary",
                    "trace_file": self.trace_path,
                    "ordered_steps": ordered_steps,
                    "aspirate_step_count": len(aspirate_steps),
                    "dispense_step_count": len(dispense_steps),
                    "cross_match_status": "pass" if ok else "fail",
                    "cross_match": str(cross_match),
                })

                orientation_tip_findings, orientation_tip_issues = analyze_orientation_and_tip_usage(
                    configs,
                    execution_plan,
                    marker_matches,
                    aspirate_steps,
                    dispense_steps,
                    cross_match,
                )

                trace_findings.extend(orientation_tip_findings)
                issues.extend(orientation_tip_issues)

                if ok and isinstance(cross_match, dict):
                    if marker_matches:
                        t_findings, t_issues = validate_trace_against_layout(
                            cross_match,
                            marker_matches,
                            marker,
                        )

                        trace_findings.extend(t_findings)
                        issues.extend(t_issues)
                    else:
                        limitations.append(
                            "Trace was parsed, but marker-based layout validation was skipped because no marker matches were found."
                        )

                    if self.tracking_var.get():
                        tracking_findings, tracking_issues = validate_tracking_source_order(
                            cross_match,
                            marker,
                        )

                        trace_findings.extend(tracking_findings)
                        issues.extend(tracking_issues)
                else:
                    issues.append(ValidationIssue(
                        "critical",
                        "trace",
                        f"Trace cross-match failed: {cross_match}",
                    ))

            except Exception as exc:
                issues.append(ValidationIssue(
                    "critical",
                    "trace",
                    f"Trace analysis failed: {type(exc).__name__}: {exc}",
                ))
        else:
            limitations.append(
                "No trace file was selected or found; trace validation was not completed."
            )

        output = create_output_model(
            lay_file=self.lay_path,
            layout_file=self.layout_path,
            trace_file=self.trace_path,
            configs=configs,
            pairs=pairs,
            execution_plan=execution_plan,
            marker_matches=marker_matches,
            layout_findings=layout_findings,
            trace_findings=trace_findings,
            issues=issues,
            limitations=limitations,
            review_script_path=review_script,
        )

        json_path = os.path.join(
            output_dir,
            f"DevPal_Lite_Output_Model_{datetime.now().strftime('%H%M%S')}.json",
        )

        Path(json_path).write_text(
            json.dumps(asdict(output), indent=2, default=str),
            encoding="utf-8",
        )

        report_format = self.report_format_var.get()

        if report_format == "pdf":
            report_path = write_pdf_report(output, output_dir)
        elif report_format == "txt":
            report_path = write_text_report(output, output_dir)
        else:
            report_path = write_docx_report(output, output_dir)

        output.report_path = report_path

        Path(json_path).write_text(
            json.dumps(asdict(output), indent=2, default=str),
            encoding="utf-8",
        )

        issue_count = len([
            issue for issue in issues
            if issue.severity in ("critical", "major")
        ])

        messagebox.showinfo(
            "DevPal Lite Complete",
            (
                f"Validation complete.\n\n"
                f"Major/Critical issues: {issue_count}\n"
                f"Report: {report_path}\n"
                f"Output Model: {json_path}\n"
                f"Trace: {self.trace_path or 'None'}"
            ),
        )

        self.status.config(text=f"Complete. Report saved: {report_path}")

    def run(self):
        self.root.mainloop()


def main():
    app = DevPalLiteApp()
    app.run()


if __name__ == "__main__":
    main()
